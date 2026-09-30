"""
Business Plan Outlook — structured content (format_version 2) and the PPTX
builder that renders it in the house style of "V7.Final_2026 Economic
Outlook.pptx": four dense management slides (Global, Indonesia, Indonesia
SWOT, Pharmaceutical Industry) with a green "Economic Outlook | <topic>"
header band, a two-line headline, blue banded tables, a budget chart and a
"<year> Business Plan" footer.

The old format was three freeform Markdown sections rendered as one text
slide each, which overflowed the slide and looked nothing like the report
management actually uses. The AI now fills the fixed structure below
instead, and the builder places every element at the reference deck's
positions.

Text fields use a tiny Markdown subset: "- " = bullet, "  - " = sub-bullet,
**bold**. "▲"/"▼" are colored green/red automatically.
"""
import copy
import io
import math
import re

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt, Emu
from lxml import etree

FORMAT_VERSION = 2

FONT = "Malgun Gothic"

# Office-theme colors as resolved in the 2026 reference deck
GREEN_BAND = RGBColor(0xC5, 0xE0, 0xB4)    # accent6 40% — header band
GREEN_LIGHT = RGBColor(0xE2, 0xF0, 0xD9)   # accent6 20% — Strength box
GREEN_MID = RGBColor(0xA9, 0xD1, 0x8E)     # accent6 60% — Strength tab
GREEN = RGBColor(0x70, 0xAD, 0x47)         # accent6
GREEN_DARK = RGBColor(0x54, 0x82, 0x35)    # accent6 75%
GREEN_TEXT = RGBColor(0x38, 0x57, 0x23)
BLUE_BAR = RGBColor(0xDE, 0xEB, 0xF7)      # accent5 20% — section title bars
BLUE_MID = RGBColor(0x9D, 0xC3, 0xE6)      # accent5 60%
BLUE_TAB = RGBColor(0xB4, 0xC7, 0xE7)      # accent1 40%
BLUE_HEAD = RGBColor(0x8F, 0xAA, 0xDC)     # accent1 60% — pharma section tabs
BLUE = RGBColor(0x44, 0x72, 0xC4)          # accent1
BLUE_BORDER = RGBColor(0x4F, 0x81, 0xBD)
BLUE_DARK = RGBColor(0x2F, 0x55, 0x97)     # accent1 75%
NAVY = RGBColor(0x00, 0x20, 0x60)
ORANGE = RGBColor(0xED, 0x7D, 0x31)        # accent2
ORANGE_LIGHT = RGBColor(0xF4, 0xB1, 0x83)  # accent2 60%
ORANGE_DARK = RGBColor(0xC5, 0x5A, 0x11)   # accent2 75%
GOLD_LIGHT = RGBColor(0xFF, 0xD9, 0x66)    # accent4 60%
GOLD_DARK = RGBColor(0xBF, 0x90, 0x00)     # accent4 75%
RED = RGBColor(0xC0, 0x00, 0x00)
RED_BRIGHT = RGBColor(0xFF, 0x00, 0x00)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x59, 0x59, 0x59)
LINE_DARK = RGBColor(0x1F, 0x4E, 0x79)

SLIDE_W_CM = 33.867
SLIDE_H_CM = 19.05


# ── Content schema ───────────────────────────────────────────────────────────

def year_labels(plan_year: int) -> dict:
    """Column labels as the 2026 deck used them: actual two years back,
    estimate for the current year, projection for the plan year."""
    return {
        "actual": str(plan_year - 2),
        "estimate": f"{plan_year - 1} (E)",
        "projection": f"{plan_year} (P)",
        "prev": str(plan_year - 1),
        "plan": str(plan_year),
    }


def default_content(plan_year: int) -> dict:
    """Empty skeleton with the reference deck's rows/labels in place — the
    AI prompt shows this shape, and the editor/exporter fall back to it for
    any field the AI left out."""
    y = year_labels(plan_year)
    growth_cols = ["Countries", y["actual"], y["estimate"], y["projection"]]
    return {
        "format_version": FORMAT_VERSION,
        "global": {
            "headline": ["", ""],
            "blocks": [
                {"title": "Global", "text": ""},
                {"title": "USA", "text": ""},
                {"title": "Eurozone", "text": ""},
                {"title": "China", "text": ""},
            ],
            "fx": {
                "title": f"Exchange Rate Forecast {y['plan']}",
                "prev_label": f"{y['prev']}\n(As of Sep AVG)",
                "next_label": f"{y['plan']}\n(Forecast)",
                "rows": [["USD/IDR", "", ""], ["USD/EUR", "", ""], ["USD/KRW", "", ""]],
                "source": "",
            },
            "growth": {
                "title": "Global Growth Rate",
                "unit": "(in %)",
                "columns": growth_cols,
                "rows": [[c, "", "", ""] for c in ("World", "USA", "Europe", "China", "Korea", "Indonesia")],
                "highlight": "Indonesia",
            },
            "oil": {
                "columns": ["Brent Crude Oil", y["actual"], y["estimate"], y["projection"]],
                "rows": [["World (USD/barrel)", "", "", ""]],
            },
            "source": "",
        },
        "indonesia": {
            "headline": ["", ""],
            "index_table": {
                "header": "Economic Index",
                "groups": [
                    {"name": "Ministry of Finance", "columns": [y["estimate"], y["projection"]]},
                    {"name": "UOB", "columns": [y["projection"]]},
                ],
                "rows": [
                    ["GDP (%)", "", "", ""],
                    ["Inflation (%)", "", "", ""],
                    ["Interest Rate Obligation 10 years (%)", "", "", ""],
                    ["Exchange Rate (USD/IDR)", "", "", ""],
                ],
            },
            "budget": {
                "headline": "",
                "title": "",
                "unit": "In Trillion IDR",
                "series": [y["prev"], y["plan"]],
                "items": [
                    {"name": n, "prev": None, "next": None, "desc": ""}
                    for n in ("Food", "Energy", "Health", "Social", "Education", "Security", "Others")
                ],
                "source": "",
            },
            "boxes": [
                {"title": "Acceleration Investment", "text": ""},
                {"title": "Increasing Export", "text": ""},
            ],
        },
        "swot": {
            "title": "Indonesia SWOT Analysis\n&\nImplication",
            "strengths": "", "weaknesses": "", "opportunities": "", "threats": "",
            "so": "", "wo": "", "st": "", "wt": "",
        },
        "pharma": {
            "global": {"headline": "", "text": ""},
            "top10": {
                "unit": "(In billion USD)",
                "year_label": y["estimate"],
                "rows": [["", ""] for _ in range(10)],
                "total": "",
                "source": "",
            },
            "asia": {"headline": "", "text": ""},
            "asia_top3": {"items": ["", "", ""], "total": "", "note": "", "source": ""},
            "indonesia": {"headline": "", "text": ""},
            "segments": [
                {"title": "Prescription Drug", "value": "", "text": ""},
                {"title": "OTC Drug", "value": "", "text": ""},
            ],
            "opportunities": "",
            "threats": "",
            "source": "",
        },
    }


def _merge(default, value):
    """Deep-merge value onto default: dict keys missing from value keep the
    default; lists/scalars from value win when present and well-typed."""
    if isinstance(default, dict):
        if not isinstance(value, dict):
            return copy.deepcopy(default)
        out = {}
        for k, dv in default.items():
            out[k] = _merge(dv, value[k]) if k in value else copy.deepcopy(dv)
        for k, v in value.items():
            if k not in out:
                out[k] = v
        return out
    if isinstance(default, list):
        return value if isinstance(value, list) else copy.deepcopy(default)
    if value is None:
        return default
    return value


def normalize_content(content: dict, plan_year: int) -> dict:
    """Fill any missing keys of a v2 content dict from the skeleton so the
    builder never has to guard against a partial AI answer or an older save."""
    return _merge(default_content(plan_year), content or {})


def is_v2(content) -> bool:
    return isinstance(content, dict) and content.get("format_version") == FORMAT_VERSION


# ── Low-level drawing helpers ────────────────────────────────────────────────

def _no_shadow(shape):
    try:
        shape.shadow.inherit = False
    except Exception:
        pass


def _rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75, shape=MSO_SHAPE.RECTANGLE):
    shp = slide.shapes.add_shape(shape, Cm(x), Cm(y), Cm(w), Cm(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    _no_shadow(shp)
    shp.text_frame.text = ""
    return shp


def _line(slide, x1, y1, x2, y2, color=LINE_DARK, width=1.0):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    ln.line.color.rgb = color
    ln.line.width = Pt(width)
    return ln


def _frame(shape, margin=0.1, anchor=MSO_ANCHOR.TOP, wrap=True):
    tf = shape.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = Cm(margin)
    tf.margin_top = tf.margin_bottom = Cm(0.05)
    tf.vertical_anchor = anchor
    return tf


def _run(paragraph, text, size, bold=False, color=BLACK, italic=False):
    r = paragraph.add_run()
    r.text = text
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = FONT
    f.color.rgb = color
    rPr = r._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", FONT)
    return r


def _rich(paragraph, text, size, bold=False, color=BLACK, italic=False):
    """Render **bold** spans and color ▲/▼ markers green/red."""
    for part in re.split(r"(\*\*[^*]+\*\*)", text or ""):
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**") and len(part) > 4
        seg = part[2:-2] if is_bold else part
        for piece in re.split(r"([▲▼])", seg):
            if not piece:
                continue
            col = GREEN if piece == "▲" else RED_BRIGHT if piece == "▼" else color
            _run(paragraph, piece, size, bold=bold or is_bold, color=col, italic=italic)


def _set_bullet(paragraph, char, font, level_indent_cm, hang_cm):
    """Attach a real PowerPoint bullet (buChar) with a hanging indent so
    wrapped lines align under the text, not under the bullet."""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Cm(level_indent_cm + hang_cm))))
    pPr.set("indent", str(-int(Cm(hang_cm))))
    for tag in ("a:buNone", "a:buChar", "a:buFont", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    buFont = etree.SubElement(pPr, qn("a:buFont"))
    buFont.set("typeface", font)
    buChar = etree.SubElement(pPr, qn("a:buChar"))
    buChar.set("char", char)


BULLETS = {
    "check": [("ü", "Wingdings"), ("Ø", "Wingdings")],   # ✓ then ➢ — Global slide
    "dot": [("•", "Arial"), ("Ø", "Wingdings")],          # • then ➢
    "arrow": [("Ø", "Wingdings"), ("§", "Wingdings")],     # ➢ then ▪
}


def parse_lines(text: str) -> list:
    """-> [(level, kind, text)] where level 0/1, kind 'bullet'|'plain'."""
    out = []
    for raw in (text or "").split("\n"):
        if not raw.strip():
            continue
        indent = len(raw) - len(raw.lstrip(" \t"))
        s = raw.strip()
        m = re.match(r"^[-*•]\s+(.*)$", s)
        if m:
            out.append((1 if indent >= 2 else 0, "bullet", m.group(1).strip()))
        else:
            out.append((0, "plain", s))
    return out


def _estimate_height_cm(lines, width_cm, size_pt, hang_cm=0.55, line_spacing=1.3):
    """Rough rendered height of parsed lines — used to pick a font size that
    fits the box, since python-pptx can't measure text. Malgun Gothic
    averages ~0.52em per Latin character."""
    char_cm = size_pt * 0.52 * 0.03528
    line_cm = size_pt * line_spacing * 0.03528
    total = 0
    for level, kind, txt in lines:
        plain = re.sub(r"\*\*", "", txt)
        avail = width_cm - (hang_cm * (level + 1) if kind == "bullet" else 0) - 0.3
        per_line = max(1, int(avail / char_cm))
        total += max(1, math.ceil(len(plain) / per_line))
    return total * line_cm + len(lines) * size_pt * 0.15 * 0.03528


def fit_size(text_or_lines, width_cm, height_cm, max_pt=12, min_pt=8):
    lines = parse_lines(text_or_lines) if isinstance(text_or_lines, str) else text_or_lines
    size = max_pt
    while size > min_pt and _estimate_height_cm(lines, width_cm, size) > height_cm:
        size -= 0.5
    return size


def _autofit_flag(tf):
    """normAutofit: if a user later edits the text in PowerPoint, it shrinks
    to fit rather than spilling out of the box."""
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    for tag in ("a:spAutoFit", "a:noAutofit", "a:normAutofit"):
        for el in bodyPr.findall(qn(tag)):
            bodyPr.remove(el)
    etree.SubElement(bodyPr, qn("a:normAutofit"))


def write_lines(tf, text, size, style="dot", color=BLACK, align=PP_ALIGN.LEFT, bold=False,
                space_after=2, first_para=True):
    """Render Markdown-subset text into a text frame with real bullets."""
    lines = parse_lines(text)
    glyphs = BULLETS.get(style, BULLETS["dot"])
    hang = 0.55
    for i, (level, kind, txt) in enumerate(lines):
        p = tf.paragraphs[0] if (first_para and i == 0) else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if kind == "bullet":
            ch, fnt = glyphs[min(level, 1)]
            _set_bullet(p, ch, fnt, level * (hang + 0.1), hang)
        _rich(p, txt, size, bold=bold, color=color)
    _autofit_flag(tf)


def text_box(slide, x, y, w, h, text, size=12, style="dot", color=BLACK, align=PP_ALIGN.LEFT,
             bold=False, fit=True, min_pt=8, anchor=MSO_ANCHOR.TOP, fill=None, line=None, line_w=0.75,
             margin=0.15):
    shp = _rect(slide, x, y, w, h, fill=fill, line=line, line_w=line_w)
    tf = _frame(shp, margin=margin, anchor=anchor)
    if fit:
        size = fit_size(text, w - 2 * margin, h - 0.1, max_pt=size, min_pt=min_pt)
    write_lines(tf, text, size, style=style, color=color, align=align, bold=bold)
    return shp


def label(slide, x, y, w, h, text, size=12, bold=False, color=BLACK, align=PP_ALIGN.LEFT,
          fill=None, line=None, italic=False, anchor=MSO_ANCHOR.MIDDLE, line_w=0.75, margin=0.15,
          shape=MSO_SHAPE.RECTANGLE):
    """Single- or multi-line plain label (newlines = paragraphs, no bullets)."""
    shp = _rect(slide, x, y, w, h, fill=fill, line=line, line_w=line_w, shape=shape)
    tf = _frame(shp, margin=margin, anchor=anchor)
    for i, ln in enumerate((text or "").split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        _rich(p, ln, size, bold=bold, color=color, italic=italic)
    return shp


def source_note(slide, x, y, w, text, align=PP_ALIGN.RIGHT):
    if not (text or "").strip():
        return
    label(slide, x, y, w, 0.45, text.strip(), size=8, italic=True, color=BLUE_DARK, align=align, margin=0.05)


def _set_cell(cell, text, size=12, bold=False, color=None, align=PP_ALIGN.CENTER, fill=None,
              indent=False):
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = True
    cell.margin_left = Cm(0.25 if not indent else 0.7)
    cell.margin_right = Cm(0.15)
    cell.margin_top = Cm(0.03)
    cell.margin_bottom = Cm(0.03)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    kwargs = {"bold": bold}
    if color is not None:
        kwargs["color"] = color
    _rich(p, str(text if text is not None else ""), size, **kwargs)
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill


def table(slide, x, y, w, col_widths, rows, row_h=0.8, header_rows=1, size=12,
          first_col_align=PP_ALIGN.LEFT, bold_rows=(), indent_rows=()):
    """Blue banded table (Medium Style 2 – Accent 1, the style the 2026 deck
    uses). rows includes the header row(s). Header text is white/bold via
    the table style; rows listed in bold_rows are bolded."""
    n_rows, n_cols = len(rows), len(col_widths)
    gf = slide.shapes.add_table(n_rows, n_cols, Cm(x), Cm(y), Cm(w), Cm(row_h * n_rows))
    tbl = gf.table
    scale = w / sum(col_widths)
    for i, cw in enumerate(col_widths):
        tbl.columns[i].width = Cm(cw * scale)
    for r in range(n_rows):
        tbl.rows[r].height = Cm(row_h)
        for c in range(n_cols):
            val = rows[r][c] if c < len(rows[r]) else ""
            is_head = r < header_rows
            _set_cell(
                tbl.cell(r, c), val, size=size,
                bold=is_head or r in bold_rows,
                color=WHITE if is_head else BLACK,
                align=(first_col_align if (c == 0 and not is_head) else PP_ALIGN.CENTER),
                indent=(c == 0 and r in indent_rows),
            )
    return gf


def header(slide, topic):
    band = _rect(slide, 0.7, 0.5, 32.5, 1.1, fill=GREEN_BAND)
    label(slide, 0.9, 0.5, 16, 1.1, "Economic Outlook", size=20, bold=True)
    label(slide, 16.9, 0.5, 16.1, 1.1, topic, size=20, bold=True, align=PP_ALIGN.RIGHT)
    return band


def footer(slide, plan_year):
    _line(slide, 0.6, 18.15, 33.1, 18.15, color=LINE_DARK, width=1.0)
    label(slide, 20, 18.2, 13.1, 0.8, f"{plan_year} Business Plan", size=12, bold=True, align=PP_ALIGN.RIGHT)


def headline(slide, lines, y=2.0, h=1.9):
    lines = [ln for ln in (lines or []) if (ln or "").strip()]
    if not lines:
        return
    shp = _rect(slide, 0.6, y, 32.5, h)
    tf = _frame(shp, anchor=MSO_ANCHOR.MIDDLE)
    longest = max(len(ln) for ln in lines)
    size = 16 if longest <= 85 else 14 if longest <= 100 else 12
    for i, ln in enumerate(lines[:2]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(4)
        _rich(p, ln.strip(), size, bold=True)


def _split_heights(texts, total_h, width_cm, fixed_per_block, min_h=1.3):
    """Share total_h between blocks in proportion to their estimated text
    height at 12pt, after reserving each block's fixed part (title bar)."""
    avail = total_h - fixed_per_block * len(texts)
    need = [max(_estimate_height_cm(parse_lines(t), width_cm, 12), min_h) for t in texts]
    s = sum(need) or 1
    return [avail * n / s for n in need]


# ── Slides ───────────────────────────────────────────────────────────────────

def _slide_global(prs, c, plan_year):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Global")
    headline(s, c.get("headline"))

    # Left: titled blocks (Global / USA / Eurozone / China)
    blocks = [b for b in (c.get("blocks") or []) if (b.get("title") or b.get("text"))][:5]
    x, w, top, bottom, bar_h, gap = 0.8, 17.1, 4.4, 17.95, 0.8, 0.15
    if blocks:
        heights = _split_heights([b.get("text", "") for b in blocks], bottom - top - gap * (len(blocks) - 1), w - 0.6, bar_h)
        y = top
        for b, h in zip(blocks, heights):
            label(s, x, y, w, bar_h, b.get("title", ""), size=12, fill=BLUE_BAR, margin=0.3)
            text_box(s, x + 0.3, y + bar_h + 0.05, w - 0.3, h - 0.05, b.get("text", ""), size=12, style="check", min_pt=9)
            y += bar_h + h + gap

    # Right top: exchange-rate forecast panel
    fx = c.get("fx") or {}
    rx, rw = 18.5, 14.7
    _rect(s, rx, 4.4, rw, 3.6, line=GREEN_DARK, line_w=2.25)
    _line(s, rx + 4.3, 4.4, rx + 4.3, 8.0, color=GREEN_DARK, width=2.25)
    label(s, rx + 0.1, 4.45, 4.1, 3.5, fx.get("title", ""), size=18, bold=True, color=GREEN_DARK, align=PP_ALIGN.CENTER)
    gx = rx + 4.5
    cols = [(gx, 2.4, PP_ALIGN.LEFT), (gx + 2.4, 3.4, PP_ALIGN.CENTER), (gx + 5.8, 1.1, PP_ALIGN.CENTER), (gx + 6.9, 3.2, PP_ALIGN.CENTER)]
    label(s, cols[1][0], 4.5, cols[1][1], 1.0, fx.get("prev_label", ""), size=11, align=PP_ALIGN.CENTER)
    label(s, cols[3][0], 4.5, cols[3][1], 1.0, fx.get("next_label", ""), size=11, align=PP_ALIGN.CENTER)
    for i, row in enumerate((fx.get("rows") or [])[:3]):
        ry = 5.55 + i * 0.8
        row = list(row) + ["", "", ""]
        vals = [row[0], row[1], "→", row[2]]
        for (cx, cw, al), v in zip(cols, vals):
            label(s, cx, ry, cw, 0.75, v, size=12, bold=(i == 0), color=RED_BRIGHT if i == 0 else RED, align=al, margin=0.05)
    source_note(s, rx, 8.05, rw, fx.get("source", ""))

    # Right middle: growth table
    g = c.get("growth") or {}
    label(s, rx, 8.55, rw, 0.8, g.get("title", ""), size=12, bold=True, color=GREEN_TEXT, fill=BLUE_BAR, align=PP_ALIGN.CENTER)
    label(s, rx + rw - 2.0, 8.55, 1.9, 0.8, g.get("unit", ""), size=10, italic=True, align=PP_ALIGN.RIGHT)
    g_rows = [g.get("columns") or []] + [list(r) for r in (g.get("rows") or [])][:8]
    row_h = 0.75 if len(g_rows) <= 7 else 0.66
    table(s, rx, 9.45, rw, [4.4, 3.4, 3.4, 3.5], g_rows, row_h=row_h, size=12,
          bold_rows=(1,), indent_rows=tuple(range(2, len(g_rows))))
    hl = (g.get("highlight") or "").strip().lower()
    for i, r in enumerate(g_rows[1:], start=1):
        if hl and r and str(r[0]).strip().lower() == hl:
            _rect(s, rx - 0.05, 9.45 + i * row_h, rw + 0.1, row_h, line=RED, line_w=2.25)
    oil_y = 9.45 + len(g_rows) * row_h + 0.2

    # Right bottom: Brent oil
    o = c.get("oil") or {}
    o_rows = [o.get("columns") or []] + [list(r) for r in (o.get("rows") or [])][:2]
    table(s, rx, oil_y, rw, [4.4, 3.4, 3.4, 3.5], o_rows, row_h=0.7, size=12)
    source_note(s, rx, oil_y + 0.7 * len(o_rows) + 0.05, rw, c.get("source", ""))

    footer(s, plan_year)


def _slide_indonesia(prs, c, plan_year):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Indonesia")
    headline(s, c.get("headline"))

    # Left top: economic index table with a grouped two-row header
    it = c.get("index_table") or {}
    groups = it.get("groups") or []
    sub_cols = [col for g in groups for col in (g.get("columns") or [])]
    n_cols = 1 + len(sub_cols)
    rows = [list(r) + [""] * n_cols for r in (it.get("rows") or [])]
    rows = [r[:n_cols] for r in rows]
    head1 = [it.get("header", "Economic Index")] + [""] * len(sub_cols)
    head2 = [""] + sub_cols
    x, y, w = 0.7, 4.0, 16.9
    widths = [7.9] + [9.0 / max(1, len(sub_cols))] * len(sub_cols)
    gf = table(s, x, y, w, widths, [head1, head2] + rows, row_h=0.85, header_rows=2, size=12)
    tbl = gf.table
    for ci, name in enumerate(sub_cols, start=1):
        _set_cell(tbl.cell(1, ci), name, bold=True, color=WHITE, fill=BLUE)
    tbl.cell(0, 0).merge(tbl.cell(1, 0))
    _set_cell(tbl.cell(0, 0), it.get("header", "Economic Index"), bold=True, color=WHITE)
    col = 1
    for g in groups:
        span = len(g.get("columns") or [])
        if span <= 0:
            continue
        if span > 1:
            tbl.cell(0, col).merge(tbl.cell(0, col + span - 1))
        _set_cell(tbl.cell(0, col), g.get("name", ""), bold=True, color=WHITE)
        col += span
    table_bottom = y + 0.85 * (2 + len(rows))

    # Left bottom: titled boxes (Acceleration Investment / Increasing Export)
    boxes = [b for b in (c.get("boxes") or []) if (b.get("title") or b.get("text"))][:3]
    if boxes:
        top, bottom, tab_h, gap = table_bottom + 0.45, 17.95, 0.8, 0.35
        heights = _split_heights([b.get("text", "") for b in boxes], bottom - top - gap * (len(boxes) - 1), w - 0.4, tab_h)
        by = top
        for b, h in zip(boxes, heights):
            label(s, x, by, 6.8, tab_h, b.get("title", ""), size=15, bold=True, fill=BLUE_TAB, margin=0.25)
            text_box(s, x + 0.1, by + tab_h, w - 0.1, h, b.get("text", ""), size=12, style="dot",
                     line=BLUE_MID, line_w=1.25, min_pt=9)
            by += tab_h + h + gap

    # Right: budget priorities + chart + category descriptions
    b = c.get("budget") or {}
    rx, rw = 17.9, 15.3
    if (b.get("headline") or "").strip():
        label(s, rx, 4.0, rw, 1.1, b["headline"], size=13 if len(b["headline"]) <= 70 else 12, bold=True, align=PP_ALIGN.CENTER)
    if (b.get("title") or "").strip():
        label(s, rx, 5.1, rw, 0.75, b["title"], size=14, bold=True, align=PP_ALIGN.CENTER)
    items = [it_ for it_ in (b.get("items") or []) if (it_.get("name") or "").strip()][:8]
    chart_items = [it_ for it_ in items if _num(it_.get("prev")) is not None or _num(it_.get("next")) is not None]
    if chart_items:
        cd = CategoryChartData()
        cd.categories = [it_["name"] for it_ in chart_items]
        series = (b.get("series") or [str(plan_year - 1), str(plan_year)]) + ["", ""]
        cd.add_series(series[0], [_num(it_.get("prev")) or 0 for it_ in chart_items])
        cd.add_series(series[1], [_num(it_.get("next")) or 0 for it_ in chart_items])
        gfc = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Cm(rx), Cm(5.85), Cm(rw), Cm(5.6), cd)
        ch = gfc.chart
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(10)
        ch.font.name = FONT
        plot = ch.plots[0]
        plot.gap_width = 80
        plot.overlap = 0
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.font.size = Pt(7)
        dl.number_format = "#,##0"
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        for sr, colr in zip(plot.series, (BLUE, ORANGE)):
            sr.format.fill.solid()
            sr.format.fill.fore_color.rgb = colr
        va = ch.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = RGBColor(0xD9, 0xD9, 0xD9)
        va.tick_labels.font.size = Pt(8)
        va.tick_labels.number_format = "#,##0"
        va.tick_labels.number_format_is_linked = False
        va.format.line.fill.background()
        if b.get("unit"):
            va.has_title = True
            va.axis_title.text_frame.text = b["unit"]
            va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
            va.axis_title.text_frame.paragraphs[0].runs[0].font.italic = True
        ca = ch.category_axis
        ca.tick_labels.font.size = Pt(10)
        ca.format.line.color.rgb = RGBColor(0xBF, 0xBF, 0xBF)
    source_note(s, rx, 11.45, rw, b.get("source", ""))

    described = [it_ for it_ in items if (it_.get("desc") or "").strip()]
    if described:
        _category_list(s, rx + 0.3, 11.95, rw - 0.3, 17.95 - 11.95, described)

    footer(s, plan_year)


def _category_list(slide, x, y, w, h, items):
    """'•  Food    : description' rows with a tab stop so every colon lines
    up and wrapped lines indent under the description, like the 2026 deck."""
    name_w = max(len(it["name"]) for it in items) * 0.23 + 0.9
    tab_cm = min(max(name_w, 2.4), 4.5)
    colon_cm = tab_cm + 0.4
    lines = [(0, "bullet", f"{it['name']} : {it['desc']}") for it in items]
    size = 12
    while size > 8 and _estimate_height_cm(lines, w - colon_cm + 0.6, size) > h:
        size -= 0.5
    shp = _rect(slide, x, y, w, h)
    tf = _frame(shp)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(1)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Cm(colon_cm))))
        pPr.set("indent", str(-int(Cm(colon_cm))))
        tabLst = etree.SubElement(pPr, qn("a:tabLst"))
        tab = etree.SubElement(tabLst, qn("a:tab"))
        tab.set("pos", str(int(Cm(tab_cm))))
        tab.set("algn", "l")
        _run(p, "•  ", size)
        _run(p, it["name"].strip(), size, bold=True)
        _run(p, "\t: ", size)
        _rich(p, it["desc"].strip(), size)
    _autofit_flag(tf)


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def _slide_swot(prs, c, plan_year):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Indonesia")

    label(s, 1.0, 2.6, 10.8, 3.6, c.get("title", ""), size=20, bold=True, align=PP_ALIGN.CENTER)

    def quadrant(x, y, w, h, fill, tab_fill, tab_text, text, tab_right=False):
        _rect(s, x, y, w, h, fill=fill, line=fill, shape=MSO_SHAPE.RECTANGLE)
        tx = x + w - 5.5 if tab_right else x
        label(s, tx, y, 5.5, 1.1, tab_text, size=20, bold=True, color=GREEN_TEXT, fill=tab_fill,
              align=PP_ALIGN.RIGHT if tab_right else PP_ALIGN.LEFT, margin=0.3)
        text_box(s, x + 0.15, y + 1.15, w - 0.3, h - 1.25, text, size=14, style="dot", min_pt=9)

    quadrant(12.0, 2.2, 10.6, 4.7, GREEN_LIGHT, GREEN_MID, "Strength", c.get("strengths", ""))
    quadrant(22.6, 2.2, 10.5, 4.7, BLUE_BAR, BLUE_MID, "Weaknesses", c.get("weaknesses", ""), tab_right=True)
    quadrant(1.3, 6.9, 10.7, 5.5, GOLD_LIGHT, GOLD_DARK, "Opportunities", c.get("opportunities", ""))
    quadrant(1.3, 12.4, 10.7, 5.5, ORANGE_LIGHT, ORANGE_DARK, "Threats", c.get("threats", ""))

    for (x, y, w, key) in ((12.0, 6.9, 10.6, "so"), (22.6, 6.9, 10.5, "wo"),
                           (12.0, 12.4, 10.6, "st"), (22.6, 12.4, 10.5, "wt")):
        text_box(s, x, y, w, 5.5, c.get(key, ""), size=14, style="dot", line=GREEN_DARK, line_w=1.0,
                 anchor=MSO_ANCHOR.MIDDLE, min_pt=9, margin=0.35)

    circ = _rect(s, 21.1, 10.9, 3.0, 3.0, fill=BLUE_TAB, shape=MSO_SHAPE.OVAL)
    for txt, lx, ly in (("SO", 21.15, 11.3), ("WO", 22.6, 11.3), ("ST", 21.15, 12.45), ("WT", 22.6, 12.45)):
        label(s, lx, ly, 1.45, 1.0, txt, size=16, bold=True, color=BLUE_DARK, align=PP_ALIGN.CENTER, margin=0)

    footer(s, plan_year)


def _section_card(slide, x, y, w, h, title, headline_text, body, tab_w=None):
    """Blue title tab over a white framed card: bold ▲ headline, then ➢ points."""
    tab_w = tab_w or (w - 1.8)
    card = _rect(slide, x, y + 0.6, w, h - 0.6, fill=WHITE, line=BLUE_BORDER, line_w=1.5)
    label(slide, x + (w - tab_w) / 2, y, tab_w, 1.05, title, size=14, bold=True, fill=BLUE_HEAD, align=PP_ALIGN.CENTER)
    inner_y, inner_h = y + 1.1, h - 1.2
    hl = (headline_text or "").strip()
    body_lines = parse_lines(body)
    size = 12
    while size > 8:
        need = _estimate_height_cm([(0, "plain", "▲ " + hl)], w - 0.4, size) + _estimate_height_cm(body_lines, w - 0.4, size)
        if need <= inner_h:
            break
        size -= 0.5
    shp = _rect(slide, x + 0.1, inner_y, w - 0.2, inner_h)
    tf = _frame(shp, anchor=MSO_ANCHOR.MIDDLE)
    first = True
    if hl:
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        if not hl.startswith(("▲", "▼")):
            _run(p, "▲ ", size, bold=True, color=GREEN)
        _rich(p, hl, size, bold=True)
        first = False
    if body_lines:
        write_lines(tf, body, size, style="arrow", align=PP_ALIGN.CENTER, first_para=first)
    _autofit_flag(tf)
    return card


def _slide_pharma(prs, c, plan_year):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Pharmaceutical Industry")

    # ── Left column ──
    lx, lw = 0.8, 16.2
    g = c.get("global") or {}
    _section_card(s, lx, 2.3, lw, 4.1, "Global Pharmaceuticals", g.get("headline"), g.get("text"), tab_w=12.9)

    t10 = c.get("top10") or {}
    rows = [list(r) + ["", ""] for r in (t10.get("rows") or [])][:10]
    rows += [["", ""]] * (10 - len(rows))
    label(s, lx + 10.8, 6.55, 5.4, 0.5, t10.get("unit", ""), size=10, italic=True, color=NAVY, align=PP_ALIGN.RIGHT)
    label(s, lx, 7.2, 3.1, 3.9, "Top 10\nGlobal\nMarket", size=14, color=BLACK, align=PP_ALIGN.CENTER,
          fill=WHITE, line=BLUE_DARK, line_w=1.25, shape=MSO_SHAPE.TEAR)
    yl = t10.get("year_label", "")
    t_rows = [["No", "Region", yl, "", "Region", yl]]
    for i in range(5):
        a, b = rows[i], rows[i + 5]
        t_rows.append([str(i + 1), a[0], a[1], str(i + 6), b[0], b[1]])
    table(s, 4.1, 7.05, 12.9, [1.1, 2.3, 2.1, 1.1, 2.3, 2.1], t_rows, row_h=0.72, size=10, first_col_align=PP_ALIGN.CENTER)
    if (t10.get("total") or "").strip():
        label(s, 4.1, 7.05 + 0.74 * 6, 12.9, 0.55, t10["total"], size=10, fill=BLUE_TAB, align=PP_ALIGN.CENTER)
    source_note(s, lx, 12.0, lw, t10.get("source", ""))

    a = c.get("asia") or {}
    _section_card(s, lx, 12.45, lw, 3.4, "Asia Pharmaceuticals", a.get("headline"), a.get("text"), tab_w=12.9)

    t3 = c.get("asia_top3") or {}
    t3_items = [i for i in (t3.get("items") or []) if (i or "").strip()][:3]
    if t3_items:
        label(s, lx, 16.05, 3.1, 1.6, "Top 3\nAsia\nMarket", size=10, align=PP_ALIGN.CENTER,
              fill=WHITE, line=BLUE_DARK, line_w=1.25, shape=MSO_SHAPE.TEAR, margin=0.05)
        text_box(s, lx + 3.3, 15.95, 5.7, 1.75, "\n".join(f"- {i}" for i in t3_items), size=11, style="dot", fit=True, min_pt=8)
        arrow = s.shapes.add_shape(MSO_SHAPE.STRIPED_RIGHT_ARROW, Cm(lx + 9.0), Cm(16.45), Cm(1.3), Cm(0.7))
        arrow.fill.solid(); arrow.fill.fore_color.rgb = BLUE; arrow.line.color.rgb = BLUE_DARK; _no_shadow(arrow)
        tot = "\n".join(v for v in ((t3.get("total") or "").strip(), (t3.get("note") or "").strip()) if v)
        label(s, lx + 10.4, 15.95, 5.8, 1.5, tot, size=10.5, align=PP_ALIGN.CENTER)
        source_note(s, lx, 17.5, lw, t3.get("source", ""))

    # ── Right column ──
    rx, rw = 17.9, 15.2
    ind = c.get("indonesia") or {}
    _section_card(s, rx, 2.3, rw, 3.9, "Indonesia Pharmaceuticals", ind.get("headline"), ind.get("text"), tab_w=13.2)

    segs = (c.get("segments") or [])[:2]
    if segs:
        if len(segs) == 2:
            _line(s, rx + rw / 2, 6.2, rx + rw / 2, 8.3, color=BLACK, width=1.0)
        seg_w = (rw - 0.8) / 2
        for i, sg in enumerate(segs):
            sx = rx + i * (seg_w + 0.8)
            box = _rect(s, sx, 6.55, seg_w, 3.5, fill=WHITE, line=GREEN, line_w=1.5)
            tf = _frame(box, anchor=MSO_ANCHOR.MIDDLE, margin=0.25)
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            _rich(p, (sg.get("title") or "").strip(), 16, bold=True)
            if (sg.get("value") or "").strip():
                p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
                _rich(p2, sg["value"].strip(), 16, bold=True)
            if (sg.get("text") or "").strip():
                write_lines(tf, sg["text"], 11.5 if len(sg["text"]) < 90 else 10, style="arrow", first_para=False)
        if len(segs) == 2:
            _line(s, rx + seg_w, 8.3, rx + seg_w + 0.8, 8.3, color=BLACK, width=1.0)

    tri_up = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Cm(rx - 0.55), Cm(10.35), Cm(1.2), Cm(1.0))
    tri_up.fill.solid(); tri_up.fill.fore_color.rgb = GREEN; tri_up.line.fill.background(); _no_shadow(tri_up)
    text_box(s, rx + 0.05, 10.35, rw - 0.05, 3.65, c.get("opportunities", ""), size=12, style="check",
             line=BLUE, line_w=1.25, min_pt=8, margin=0.75)
    tri_dn = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Cm(rx - 0.55), Cm(14.15), Cm(1.2), Cm(1.0))
    tri_dn.rotation = 180
    tri_dn.fill.solid(); tri_dn.fill.fore_color.rgb = RED_BRIGHT; tri_dn.line.fill.background(); _no_shadow(tri_dn)
    text_box(s, rx + 0.05, 14.15, rw - 0.05, 3.35, c.get("threats", ""), size=12, style="check",
             line=RED_BRIGHT, line_w=1.25, min_pt=8, margin=0.75)
    source_note(s, rx, 17.5, rw, c.get("source", ""))

    footer(s, plan_year)


def build_outlook_pptx(content: dict, plan_year: int) -> bytes:
    c = normalize_content(content, plan_year)
    prs = Presentation()
    prs.slide_width = Cm(SLIDE_W_CM)
    prs.slide_height = Cm(SLIDE_H_CM)
    _slide_global(prs, c["global"], plan_year)
    _slide_indonesia(prs, c["indonesia"], plan_year)
    _slide_swot(prs, c["swot"], plan_year)
    _slide_pharma(prs, c["pharma"], plan_year)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ── AI prompt ────────────────────────────────────────────────────────────────

def ai_schema_example(plan_year: int) -> str:
    """The JSON shape the AI must return, pre-filled with this plan year's
    column labels and illustrative (NOT real) values, so the model sees
    the exact structure, the wording style and the density expected."""
    import json
    y = year_labels(plan_year)
    P, E = y["plan"], y["prev"]
    ex = default_content(plan_year)
    ex["global"]["headline"] = [
        f"“The GDP is projected to increase 0.1%p to 3.1% in {P}”",
        f"Global inflation is expected to fall from 4.2% in {E} to 3.6% in {P}",
    ]
    ex["global"]["blocks"] = [
        {"title": "Global", "text": "- Uncertainty of US reciprocal tariff rates could weaken the growth\n- Geopolitical tension disrupt global supply chains :\n  - US - China strategic rivalry across trade, technology and security\n  - Middle East tension"},
        {"title": "USA", "text": f"- The Fed rate target is projected to decrease from 4.00% in {E} to 3.75% in {P}\n- Growth is projected to slightly increase in {P} at 2.0% from 1.9% in {E}"},
        {"title": "Eurozone", "text": f"- Growth is expected to accelerate in {P} (1.2%) driven by ..."},
        {"title": "China", "text": "- Aging population is impacting the pension system, healthcare services and labor supply"},
    ]
    ex["global"]["fx"]["rows"] = [["USD/IDR", "16,404", "16,500"], ["USD/EUR", "0.89", "0.86"], ["USD/KRW", "1,441", "1,330"]]
    ex["global"]["fx"]["source"] = f"(Source : RAPBN {P}, ECB and UOB Outlook)"
    ex["global"]["growth"]["rows"] = [["World", "3.3", "3.0", "3.1"], ["USA", "2.8", "1.9", "2.0"], ["Europe", "0.9", "1.0", "1.2"], ["China", "5.0", "4.8", "4.2"], ["Korea", "2.0", "0.8", "1.8"], ["Indonesia", "5.0", "4.8", "4.8"]]
    ex["global"]["oil"]["rows"] = [["World (USD/barrel)", "81.26", "68.18", "64.33"]]
    ex["global"]["source"] = f"(Source : IMF - July {E})"
    ex["indonesia"]["headline"] = [
        f"“The GDP is projected to increase 0.5%p to 5.4% in {P}”",
        f"Inflation is expected to increase from 2.4% in {E} to 2.5% in {P}",
    ]
    ex["indonesia"]["index_table"]["rows"] = [["GDP (%)", "4.9", "5.4", "5.2"], ["Inflation (%)", "2.4", "2.5", "2.5"], ["Interest Rate Obligation 10 years (%)", "7.0", "6.9", "-"], ["Exchange Rate (USD/IDR)", "16,550", "16,500", "16,333"]]
    ex["indonesia"]["budget"].update({
        "headline": f"“RAPBN {P}, Prioritize in Food Security, Energy, Education, and Health”",
        "title": "Budget Allocation IDR 3,787 trillion",
        "items": [
            {"name": "Food", "prev": 155, "next": 164, "desc": "Increasing productivity, price stability and welfare of farmers"},
            {"name": "Health", "prev": 211, "next": 244, "desc": "Improve the quality of the healthcare system"},
            {"name": "…", "prev": 0, "next": 0, "desc": "…"},
        ],
        "source": f"(Source : RAPBN {P})",
    })
    ex["indonesia"]["boxes"] = [
        {"title": "Acceleration Investment", "text": "- **DANANTARA** : State-Owned Enterprises (SOE) super holding company\n- **Main function** :\n  - Indonesia investment acceleration\n  - Enhancing competitiveness"},
        {"title": "Increasing Export", "text": "- Negotiation US tariff from 32% to 19%, more competitive than other ASEAN countries\n- Comprehensive Economic Partnership Agreement (CEPA) with some countries"},
    ]
    ex["swot"].update({
        "strengths": f"- **Resilient Economic Growth** : Projected growth 5.4% in {P}, supported by public spending\n- **US Tariff** : Successfully negotiated from 32% to 19%",
        "weaknesses": "- **Budget Deficit Concerns** : ...\n- **Inequality** : ...",
        "opportunities": "- **Infrastructure Development** : ...\n- **Increasing Foreign Investment** : ...",
        "threats": "- **Global Economic Slowdown** : ...\n- **Market Competition** : ...",
        "so": "- Implement the downstream strategy to increase export and reduce import reliance\n- ...",
        "wo": "- ...\n- ...", "st": "- ...\n- ...", "wt": "- ...\n- ...",
    })
    ph = ex["pharma"]
    ph["global"] = {"headline": f"Expected increase in sales from USD 1.9 trillion in {E} to USD 2.0 trillion in {P} (▲ 5%)", "text": "- Technological integration, particularly AI in diagnostics and treatment, is accelerating growth"}
    ph["top10"]["rows"] = [["US", "491"], ["China", "285"], ["Japan", "100"], ["Germany", "89"], ["France", "57"], ["UK", "50"], ["Italy", "41"], ["Canada", "39"], ["India", "36"], ["Spain", "35"]]
    ph["top10"]["total"] = "Total USD 1.2 trillion (market share 64% from total sales)"
    ph["top10"]["source"] = f"(Source : BMI Global - Q3 {E})"
    ph["asia"] = {"headline": f"Expected to increase sales from USD 425 billion in {E} to USD 459 billion in {P} (▲ 8%)", "text": "- China's role in global clinical research will expand"}
    ph["asia_top3"] = {"items": ["China USD 285 billion", "Japan USD 100 billion", "India USD 36 billion"], "total": "Total USD 421 billion", "note": "(market share 99% from total sales)", "source": f"(Source : BMI Global - Q3 {E})"}
    ph["indonesia"] = {"headline": f"Expected to increase sales from IDR 122 trillion in {E} to IDR 133 trillion in {P} (▲ 9%)", "text": "- The government is increasing investment in healthcare infrastructure while implementing cost control"}
    ph["segments"] = [
        {"title": "Prescription Drug", "value": "IDR 104 trillion", "text": "- Patented drug : IDR 13 trillion\n- Generic drug : IDR 91 trillion"},
        {"title": "OTC Drug", "value": "IDR 29 trillion", "text": ""},
    ]
    ph["opportunities"] = "- Expansion of BPJS increases medical demand, especially for generic drugs\n- ...\n- ..."
    ph["threats"] = "- Vulnerable to global supply chain disruption as Indonesia relies on imported raw materials\n- ...\n- ..."
    ph["source"] = f"(Source : BMI Indonesia - August {E})"
    return json.dumps(ex, ensure_ascii=False, indent=1)
