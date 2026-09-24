"""
Markdown tables -> .xlsx for the Open WebUI "Export Excel" action.

Files are kept for an hour under an unguessable token and served without
auth: the browser following the link cannot present the service key, and
the content is a table the same user has just been shown in chat. Expired
files are removed whenever a new one is written.
"""
import re
import secrets
import tempfile
import time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

EXPORT_DIR = Path(tempfile.gettempdir()) / "ebs_exports"
TTL_SECONDS = 3600
_TOKEN_RE = re.compile(r"^[A-Za-z0-9_-]{20,64}$")
_SEPARATOR = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")
# 1.250.000.000 / 1.250.000,50 / -12.345 (Indonesian grouping)
_ID_NUMBER = re.compile(r"^-?\d{1,3}(\.\d{3})+(,\d+)?$|^-?\d+(,\d+)?$")


def _cells(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip().replace("**", "") for c in line.split("|")]


def parse_tables(markdown: str) -> list[list[list[str]]]:
    lines = markdown.splitlines()
    tables, i = [], 0
    while i < len(lines) - 1:
        if "|" in lines[i] and _SEPARATOR.match(lines[i + 1] or ""):
            rows = [_cells(lines[i])]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(_cells(lines[i]))
                i += 1
            tables.append(rows)
        else:
            i += 1
    return tables


def _value(text: str):
    t = text.strip()
    raw = re.sub(r"^(Rp\.?|IDR|USD|EUR)\s*", "", t, flags=re.I).strip()
    if _ID_NUMBER.match(raw):
        try:
            return float(raw.replace(".", "").replace(",", "."))
        except ValueError:
            return t
    return t


def _cleanup():
    now = time.time()
    for f in EXPORT_DIR.glob("*.xlsx"):
        if now - f.stat().st_mtime > TTL_SECONDS:
            f.unlink(missing_ok=True)


def build_xlsx(markdown: str, title: str) -> tuple[str, int]:
    tables = parse_tables(markdown)
    if not tables:
        raise ValueError("Tidak ada tabel markdown di jawaban.")
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    _cleanup()

    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for n, rows in enumerate(tables, start=1):
        ws = wb.create_sheet(f"Tabel {n}")
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append([])
        for r, row in enumerate(rows):
            ws.append(row if r == 0 else [_value(c) for c in row])
        for cell in ws[3]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for col in range(1, max(len(r) for r in rows) + 1):
            width = max(len(str(r[col - 1])) if col - 1 < len(r) else 0 for r in rows)
            ws.column_dimensions[get_column_letter(col)].width = min(max(10, width + 2), 60)
            for cell in ws.iter_rows(min_row=4, min_col=col, max_col=col):
                if isinstance(cell[0].value, float):
                    cell[0].number_format = "#,##0.##"
        ws.freeze_panes = "A4"

    token = secrets.token_urlsafe(24)
    wb.save(EXPORT_DIR / f"{token}.xlsx")
    return token, len(tables)


def export_path(token: str) -> Path | None:
    if not _TOKEN_RE.match(token or ""):
        return None
    path = EXPORT_DIR / f"{token}.xlsx"
    if not path.exists() or time.time() - path.stat().st_mtime > TTL_SECONDS:
        return None
    return path
