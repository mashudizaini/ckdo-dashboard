"""
Purchase Plan (Finished Good) Service
Parallel to purchase_plan_service.py's PurchasePlanService, but parses a
genuinely different Excel template ("Purchase Plan FG - 2026.xlsx") — see
that file's docstring / the "Purchase Plan (Finished Good)" plan for the
full column-layout comparison against the Material template.
"""
import io
from datetime import datetime
from typing import Optional
from sqlalchemy import select, delete
from sqlalchemy.ext.asyncio import AsyncSession
from app.models.purchase_plan_fg import PurchasePlanFinishedGood
import structlog

logger = structlog.get_logger()

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

FG_HEADERS = [
    "No", "Item Code", "Name", "Stock",
    *[f"Usage {m}" for m in MONTHS], "Usage 1H Next Year", "Usage Total",
    *[f"Order {m}" for m in MONTHS], "Order 1H Next Year", "Order Total",
    "Unit Price (Orig)", "Unit Price (IDR)", "Total Price (Rp)",
]

PP_CATEGORIES = ["Summary", "Local", "CMO", "Export"]


def _num(v):
    """Blank template cells sometimes hold stray non-numeric placeholders —
    treat anything that isn't a real number as "no value yet"."""
    if isinstance(v, (int, float)):
        return v
    try:
        return float(v) if v not in (None, "") else 0
    except (TypeError, ValueError):
        return 0


def _default_content():
    return {
        "meta": {"type": "", "department": "", "team_code": "", "team_name": "", "exchange_rate": 0},
        "headers": FG_HEADERS,
        "items": [],
    }


class PurchasePlanFGService:

    async def list_purchase_plans(
        self,
        db: AsyncSession,
        plan_year: Optional[int] = None,
        department: Optional[str] = None,
        team_code: Optional[str] = None,
        plan_category: Optional[str] = None,
    ) -> dict:
        q = select(PurchasePlanFinishedGood).order_by(
            PurchasePlanFinishedGood.plan_year.desc(), PurchasePlanFinishedGood.department, PurchasePlanFinishedGood.team_code
        )
        if plan_year:
            q = q.where(PurchasePlanFinishedGood.plan_year == plan_year)
        if department:
            q = q.where(PurchasePlanFinishedGood.department == department)
        if team_code:
            q = q.where(PurchasePlanFinishedGood.team_code == team_code)
        if plan_category:
            q = q.where(PurchasePlanFinishedGood.plan_category == plan_category)
        result = await db.execute(q)
        rows = result.scalars().all()
        return {"success": True, "count": len(rows), "data": [self._to_dict(r) for r in rows]}

    async def get_purchase_plan(self, db: AsyncSession, plan_id: int) -> dict:
        row = await db.get(PurchasePlanFinishedGood, plan_id)
        if not row:
            return {"success": False, "error": "Not found"}
        return {"success": True, "data": self._to_dict(row)}

    async def upsert_purchase_plan(self, db: AsyncSession, payload: dict, username: str) -> dict:
        plan_id = payload.get("id")
        row = None
        if plan_id:
            row = await db.get(PurchasePlanFinishedGood, plan_id)
        if not row:
            q = select(PurchasePlanFinishedGood).where(
                PurchasePlanFinishedGood.plan_year     == payload.get("plan_year"),
                PurchasePlanFinishedGood.plan_category == payload.get("plan_category", "Local"),
                PurchasePlanFinishedGood.department    == payload.get("department", ""),
                PurchasePlanFinishedGood.team_code     == str(payload.get("team_code", "")),
            )
            result = await db.execute(q)
            row = result.scalar_one_or_none()
        if row:
            row.content    = payload.get("content", row.content)
            row.status     = payload.get("status",  row.status)
            row.updated_at = datetime.utcnow()
        else:
            row = PurchasePlanFinishedGood(
                plan_year     = payload.get("plan_year", datetime.now().year),
                plan_category = payload.get("plan_category", "Local"),
                department    = payload.get("department", ""),
                team_code     = str(payload.get("team_code", "")),
                team_name     = payload.get("team_name", ""),
                content       = payload.get("content", _default_content()),
                status        = payload.get("status", "draft"),
                created_by    = username,
            )
            db.add(row)
        await db.flush()
        await db.refresh(row)
        return {"success": True, "data": self._to_dict(row)}

    async def delete_purchase_plan(self, db: AsyncSession, plan_id: int) -> dict:
        row = await db.get(PurchasePlanFinishedGood, plan_id)
        if not row:
            return {"success": False, "error": "Not found"}
        await db.execute(delete(PurchasePlanFinishedGood).where(PurchasePlanFinishedGood.id == plan_id))
        return {"success": True, "message": f"Deleted purchase plan (FG) #{plan_id}"}

    def _read_sheet_items(self, ws) -> list:
        """Each real item spans two physical rows: a "Usage" row (columns
        A-D = No/Item Code/Name/Stock, F-Q = Jan-Dec usage qty, R = 1H next
        year, S = Total, T-V = Unit Price orig/IDR + Total Price) directly
        followed by an "Order" row (blank A/C, F-Q/R/S = order qty/1H/total
        only). A "Total" grand-total row pair (A = literal "Total") follows
        the real items and is explicitly skipped — unlike Material's sheets,
        it isn't naturally caught by the "blank No+Name" skip check since
        its own A/C cells ARE filled (with "Total" / blank respectively).
        Header at rows 13-14, data from row 15."""
        items = []
        r = 15
        while r <= ws.max_row:
            no = ws.cell(row=r, column=1).value
            name = ws.cell(row=r, column=3).value
            if no is None and name is None:
                r += 1
                continue
            if str(no or "").strip().lower() == "total":
                r += 2  # this grand-total row + its paired Order-totals row
                continue

            usage_months = [_num(ws.cell(row=r, column=c).value) for c in range(6, 18)]   # F-Q
            item = {
                "no": no,
                "item_code": ws.cell(row=r, column=2).value or "",
                "name": str(name or ""),
                "stock": _num(ws.cell(row=r, column=4).value) or None,
                "usage": usage_months,
                "usage_1h_next": _num(ws.cell(row=r, column=18).value),   # R
                "usage_total": _num(ws.cell(row=r, column=19).value),    # S
                "unit_price_orig": _num(ws.cell(row=r, column=20).value),  # T
                "unit_price_idr": _num(ws.cell(row=r, column=21).value),   # U
                "total_price": _num(ws.cell(row=r, column=22).value),      # V
                "order": [0] * 12,
                "order_1h_next": 0,
                "order_total": 0,
            }
            # The next row, if it's the paired "Order" row (no No/Name of
            # its own), holds order qty.
            if r + 1 <= ws.max_row:
                next_no = ws.cell(row=r + 1, column=1).value
                next_name = ws.cell(row=r + 1, column=3).value
                if next_no is None and next_name is None:
                    item["order"] = [_num(ws.cell(row=r + 1, column=c).value) for c in range(6, 18)]
                    item["order_1h_next"] = _num(ws.cell(row=r + 1, column=18).value)
                    item["order_total"] = _num(ws.cell(row=r + 1, column=19).value)
                    r += 1
            items.append(item)
            r += 1
        return items

    async def import_excel(self, db: AsyncSession, file_bytes: bytes, plan_year: int, username: str) -> dict:
        """Parse an uploaded Excel matching the "Purchase Plan FG -
        2026.xlsx" template — one plan per data sheet, each with its own
        Type/Department/Team meta and item list.

        Sheet-match signature is A6 (not A1 like the Material template) —
        verified live that a real "_Export" copy of this template has a
        broken '#REF!' in A1 (an external-link artifact), but A6's
        " [ Type ]" label survives intact. Meta value columns are also
        shifted vs Material: C (not D) for type/department/team-code, T
        (not X) for the exchange rate — verified live against sumber/
        Purchase Plan FG - 2026.xlsx.
        """
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
        except Exception as e:
            return {"success": False, "error": f"Could not read Excel file: {e}"}

        imported = []
        for ws in wb.worksheets:
            if str(ws.cell(row=6, column=1).value or "").strip().lower() != "[ type ]":
                continue

            meta = {
                "type":          ws.cell(row=6,  column=3).value or "",
                "department":    ws.cell(row=9,  column=3).value or "",
                "team_code":     ws.cell(row=11, column=3).value or "",
                "team_name":     str(ws.cell(row=11, column=4).value or "").lstrip("/ ").strip(),
                "exchange_rate": _num(ws.cell(row=12, column=20).value),  # T12
            }
            items = self._read_sheet_items(ws)
            if not items:
                continue

            category = self._infer_category(meta["type"], ws.title)
            content = {"meta": meta, "headers": FG_HEADERS, "items": items}
            payload = {
                "plan_year":     plan_year,
                "plan_category": category,
                "department":    meta["department"],
                "team_code":     str(meta["team_code"]),
                "team_name":     meta["team_name"],
                "content":       content,
                "status":        "draft",
            }
            result = await self.upsert_purchase_plan(db, payload, username)
            if result["success"]:
                imported.append({"category": category, "items": len(items), "id": result["data"]["id"]})

        if not imported:
            return {"success": False, "error": "No recognizable data sheets found — none of the sheets in this "
                                                 "file match the Purchase Plan (FG) template layout (expected "
                                                 "'[ Type ]' in cell A6 of each data sheet)."}
        return {"success": True, "imported": imported}

    def _infer_category(self, type_text: str, sheet_name: str) -> str:
        for text in (type_text, sheet_name):
            low = str(text or "").lower()
            for category in PP_CATEGORIES:
                if category.lower() in low:
                    return category
            if "total" in low:
                return "Summary"
        return "Local"

    def _to_dict(self, row: PurchasePlanFinishedGood) -> dict:
        return {
            "id":            row.id,
            "plan_year":     row.plan_year,
            "plan_category": row.plan_category,
            "department":    row.department,
            "team_code":     row.team_code,
            "team_name":     row.team_name,
            "content":       row.content,
            "status":        row.status,
            "created_by":    row.created_by,
            "created_at":    row.created_at.isoformat() if row.created_at else None,
            "updated_at":    row.updated_at.isoformat() if row.updated_at else None,
        }
