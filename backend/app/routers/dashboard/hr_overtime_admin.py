"""
Digital Overtime Management System — HRGA administration.
Route prefix: /api/v1/dashboard/hr/overtime/admin

Everything here is gated by overtime_service.require_admin (Keycloak
hr_staff/admin, or anyone HRGA marks `hrga_admin` in the approval matrix).
Covers site map D (System Setting) plus the two HRGA-only report screens,
C.9 Overtime Calculation and C.10 Overtime Monitoring & Report.
"""
import io
from datetime import date, timedelta
from decimal import Decimal
from typing import Optional

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.models.employee import Employee
from app.models.overtime import (
    PLAN_APPROVED,
    ROLE_DEPT_HEAD, ROLE_HRGA_ADMIN, ROLE_MEMBER, ROLE_TEAM_HEAD,
    RZ_APPROVED, RZ_DH_APPROVED, RZ_SUBMITTED, RZ_TH_APPROVED,
    OvertimeApprovalMatrix, OvertimeCutoffConfig,
    OvertimeRateConfig, OvertimeRequest, OvertimeRule,
)
from app.services import overtime_service as svc
from app.services.overtime_service import OvertimeActor, require_admin

router = APIRouter()

_ROLE_LEVELS = (ROLE_MEMBER, ROLE_TEAM_HEAD, ROLE_DEPT_HEAD, ROLE_HRGA_ADMIN)


# ─────────────────────────────────────────
# SCHEMAS
# ─────────────────────────────────────────

class MatrixIn(BaseModel):
    employee_id: str
    role_level: str = ROLE_MEMBER
    team_head_id: Optional[str] = None
    dept_head_id: Optional[str] = None
    is_active: bool = True


class MatrixBulkIn(BaseModel):
    employee_ids: list[str]
    role_level: Optional[str] = None
    team_head_id: Optional[str] = None
    dept_head_id: Optional[str] = None


class RuleIn(BaseModel):
    title: str
    content: Optional[str] = None
    seq: int = 0
    is_active: bool = True


class CutoffIn(BaseModel):
    name: str = "Default"
    start_day: int
    end_day: int


class NotificationSettingIn(BaseModel):
    event_key: str
    in_app: bool = True
    email: bool = False
    is_active: bool = True


class RateIn(BaseModel):
    grade: str
    hourly_rate: float


# ─────────────────────────────────────────
# PERIOD HELPERS
# ─────────────────────────────────────────

async def _resolve_window(
    db: AsyncSession, period: Optional[str], date_from: Optional[str], date_to: Optional[str]
) -> tuple[date, date, str]:
    """An explicit date range wins; otherwise the cut-off period named by
    `period` (YYYY-MM of the closing month); otherwise the period that is
    open today."""
    if date_from and date_to:
        try:
            return date.fromisoformat(date_from), date.fromisoformat(date_to), f"{date_from} - {date_to}"
        except ValueError:
            raise HTTPException(400, "date_from / date_to must be YYYY-MM-DD")

    cfg = await svc.get_active_cutoff(db)
    if not period:
        today = date.today()
        # Past the closing day, today already belongs to the next period.
        key_month = today if today.day <= cfg.end_day else svc.add_months(today, 1)
        period = f"{key_month.year:04d}-{key_month.month:02d}"
    p = svc.cutoff_period_for_key(period, cfg.start_day, cfg.end_day)
    return p.start, p.end, p.label


@router.get("/periods")
async def list_periods(
    count: int = Query(18, le=48),
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    cfg = await svc.get_active_cutoff(db)
    today = date.today()
    anchor = today if today.day <= cfg.end_day else svc.add_months(today, 1)
    out = []
    for i in range(count):
        m = svc.add_months(date(anchor.year, anchor.month, 1), -i)
        p = svc.cutoff_period_for_key(f"{m.year:04d}-{m.month:02d}", cfg.start_day, cfg.end_day)
        out.append({"key": p.key, "label": p.label, "start": p.start.isoformat(), "end": p.end.isoformat()})
    return {"cutoff": {"start_day": cfg.start_day, "end_day": cfg.end_day, "name": cfg.name}, "periods": out}


# ─────────────────────────────────────────
# C.9 — OVERTIME CALCULATION
# ─────────────────────────────────────────

async def _calculation_rows(
    db: AsyncSession, start: date, end: date,
    department: Optional[str], team: Optional[str], employee_id: Optional[str],
) -> list[OvertimeRequest]:
    conditions = [
        OvertimeRequest.ot_date >= start,
        OvertimeRequest.ot_date <= end,
        OvertimeRequest.rz_status == RZ_APPROVED,
    ]
    if department:
        conditions.append(OvertimeRequest.department == department)
    if team:
        conditions.append(OvertimeRequest.team == team)
    if employee_id:
        conditions.append(OvertimeRequest.employee_id == employee_id)
    result = await db.execute(
        select(OvertimeRequest).where(*conditions)
        .order_by(OvertimeRequest.department, OvertimeRequest.team, OvertimeRequest.employee_name, OvertimeRequest.ot_date)
    )
    return list(result.scalars().all())


def _group(rows: list[OvertimeRequest], key_fn, label: str) -> list[dict]:
    buckets: dict[str, dict] = {}
    for r in rows:
        key = key_fn(r) or "(unassigned)"
        b = buckets.setdefault(key, {
            label: key, "weekday_minutes": 0, "weekend_minutes": 0, "total_minutes": 0,
            "overtime_index": Decimal(0), "estimated_cost": Decimal(0), "entries": 0, "employees": set(),
        })
        minutes = r.payable_minutes or 0
        b["weekend_minutes" if r.overtime_type == "weekend" else "weekday_minutes"] += minutes
        b["total_minutes"] += minutes
        b["overtime_index"] += Decimal(r.overtime_index or 0)
        b["estimated_cost"] += Decimal(r.estimated_cost or 0)
        b["entries"] += 1
        b["employees"].add(r.employee_id)

    out = []
    for b in buckets.values():
        out.append({
            label: b[label],
            "weekday_hours": svc.fmt_hours(b["weekday_minutes"]),
            "weekend_hours": svc.fmt_hours(b["weekend_minutes"]),
            "total_hours": svc.fmt_hours(b["total_minutes"]),
            "overtime_index": float(b["overtime_index"]),
            "estimated_cost": float(b["estimated_cost"]),
            "entries": b["entries"],
            "employee_count": len(b["employees"]),
        })
    return sorted(out, key=lambda x: x["total_hours"], reverse=True)


@router.get("/calculation")
async def calculation(
    period: Optional[str] = Query(None, description="Cut-off period key, YYYY-MM of the closing month"),
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    department: Optional[str] = None,
    team: Optional[str] = None,
    employee_id: Optional[str] = None,
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    """Payable overtime for a cut-off period — the figure that goes to
    payroll. Only HRGA-approved realizations are counted, at their frozen
    payable_minutes; anything still in the approval chain is reported
    separately as `pending` so nobody closes a period with work in flight."""
    start, end, label = await _resolve_window(db, period, date_from, date_to)
    rows = await _calculation_rows(db, start, end, department, team, employee_id)

    pending = (await db.execute(
        select(func.count(OvertimeRequest.id)).where(
            OvertimeRequest.ot_date >= start, OvertimeRequest.ot_date <= end,
            OvertimeRequest.rz_status.in_([RZ_SUBMITTED, RZ_TH_APPROVED, RZ_DH_APPROVED]))
    )).scalar() or 0

    detail = [{
        "id": r.id, "request_no": r.request_no,
        "employee_id": r.employee_id, "employee_name": r.employee_name,
        "department": r.department, "team": r.team,
        "ot_date": r.ot_date.isoformat(), "overtime_type": r.overtime_type,
        "work_category": r.work_category, "task_description": r.task_description,
        "plan_start": r.plan_start, "plan_finish": r.plan_finish,
        "actual_start": r.actual_start, "actual_finish": r.actual_finish,
        "planned_hours": svc.fmt_hours(r.planned_minutes),
        "actual_hours": svc.fmt_hours(r.actual_minutes),
        "payable_hours": svc.fmt_hours(r.payable_minutes),
        "gap_hours": svc.fmt_hours(r.gap_minutes),
        "overtime_index": float(r.overtime_index or 0),
        "hourly_rate": float(r.hourly_rate or 0),
        "estimated_cost": float(r.estimated_cost or 0),
        "approved_at": r.hr_decision_at.isoformat() if r.hr_decision_at else None,
    } for r in rows]

    total_minutes = sum(r.payable_minutes or 0 for r in rows)
    return {
        "period": {"label": label, "start": start.isoformat(), "end": end.isoformat(), "key": period},
        "totals": {
            "total_hours": svc.fmt_hours(total_minutes),
            "weekday_hours": svc.fmt_hours(sum(r.payable_minutes or 0 for r in rows if r.overtime_type == "weekday")),
            "weekend_hours": svc.fmt_hours(sum(r.payable_minutes or 0 for r in rows if r.overtime_type == "weekend")),
            "overtime_index": float(sum(Decimal(r.overtime_index or 0) for r in rows)),
            "estimated_cost": float(sum(Decimal(r.estimated_cost or 0) for r in rows)),
            "entries": len(rows),
            "employee_count": len({r.employee_id for r in rows}),
            "pending_in_period": pending,
        },
        "detail": detail,
        "by_department": _group(rows, lambda r: r.department, "department"),
        "by_team": _group(rows, lambda r: r.team, "team"),
        "by_employee": _group(rows, lambda r: f"{r.employee_id}|{r.employee_name}", "employee"),
    }


@router.get("/calculation/export")
async def export_calculation(
    period: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    department: Optional[str] = None,
    team: Optional[str] = None,
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    start, end, label = await _resolve_window(db, period, date_from, date_to)
    rows = await _calculation_rows(db, start, end, department, team, None)

    wb = openpyxl.Workbook()
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill("solid", fgColor="1D4ED8")

    def _write(ws, headers, data_rows, widths):
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in data_rows:
            ws.append(row)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Detail"
    _write(ws,
        ["Request No", "NIK", "Name", "Department", "Team", "Date", "Category", "Work Type",
         "Task / Target Result", "Plan Start", "Plan Finish", "Planned (h)",
         "Actual Start", "Actual Finish", "Actual (h)", "Gap (h)", "Payable (h)",
         "OT Index (h)", "Hourly Rate", "Estimated Cost"],
        [[r.request_no, r.employee_id, r.employee_name, r.department, r.team,
          r.ot_date.strftime("%d-%b-%Y"), r.overtime_type, r.work_category, r.task_description,
          r.plan_start, r.plan_finish, svc.fmt_hours(r.planned_minutes),
          r.actual_start, r.actual_finish, svc.fmt_hours(r.actual_minutes),
          svc.fmt_hours(r.gap_minutes), svc.fmt_hours(r.payable_minutes),
          float(r.overtime_index or 0), float(r.hourly_rate or 0), float(r.estimated_cost or 0)]
         for r in rows],
        [16, 12, 26, 20, 20, 13, 11, 11, 46, 11, 11, 11, 11, 11, 11, 10, 11, 11, 14, 16])

    for sheet_title, grouped, key in (
        ("By Department", _group(rows, lambda r: r.department, "department"), "department"),
        ("By Team", _group(rows, lambda r: r.team, "team"), "team"),
    ):
        s = wb.create_sheet(sheet_title)
        _write(s, [key.title(), "Employees", "Entries", "Weekday (h)", "Weekend (h)", "Total (h)", "OT Index (h)", "Estimated Cost"],
               [[g[key], g["employee_count"], g["entries"], g["weekday_hours"], g["weekend_hours"],
                 g["total_hours"], g["overtime_index"], g["estimated_cost"]] for g in grouped],
               [28, 12, 10, 13, 13, 12, 13, 16])

    s = wb.create_sheet("By Employee")
    emp_rows = _group(rows, lambda r: f"{r.employee_id}|{r.employee_name}", "employee")
    _write(s, ["NIK", "Name", "Entries", "Weekday (h)", "Weekend (h)", "Total (h)", "OT Index (h)", "Estimated Cost"],
           [[g["employee"].split("|")[0], g["employee"].split("|")[-1], g["entries"], g["weekday_hours"],
             g["weekend_hours"], g["total_hours"], g["overtime_index"], g["estimated_cost"]] for g in emp_rows],
           [12, 26, 10, 13, 13, 12, 13, 16])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"Overtime_Calculation_{start:%Y%m%d}_{end:%Y%m%d}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ─────────────────────────────────────────
# C.10 — MONITORING & REPORT
# ─────────────────────────────────────────

@router.get("/monitoring")
async def monitoring(
    year: int = Query(...),
    compare_year: Optional[int] = None,
    department: Optional[str] = None,
    team: Optional[str] = None,
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    """Everything the Monitoring tab draws: weekday vs weekend by year and
    by month, planned vs realized trend per department/team, and cost.

    Realized hours here are payable_minutes on HRGA-approved rows; planned
    hours are planned_minutes on approved *orders*, including those whose
    realization never came in — that difference is exactly what the
    Exception Report is meant to surface, so the two must not be filtered
    to the same population."""
    compare_year = compare_year or (year - 1)

    def _base(conditions):
        if department:
            conditions.append(OvertimeRequest.department == department)
        if team:
            conditions.append(OvertimeRequest.team == team)
        return conditions

    async def _fetch(y: int) -> list[OvertimeRequest]:
        result = await db.execute(select(OvertimeRequest).where(*_base([
            OvertimeRequest.ot_date >= date(y, 1, 1),
            OvertimeRequest.ot_date <= date(y, 12, 31),
            OvertimeRequest.plan_status == PLAN_APPROVED,
        ])))
        return list(result.scalars().all())

    rows_cur, rows_cmp = await _fetch(year), await _fetch(compare_year)

    def _realized(r: OvertimeRequest) -> int:
        return (r.payable_minutes or 0) if r.rz_status == RZ_APPROVED else 0

    def _yearly(rows: list[OvertimeRequest], y: int) -> dict:
        return {
            "year": y,
            "weekday_hours": svc.fmt_hours(sum(_realized(r) for r in rows if r.overtime_type == "weekday")),
            "weekend_hours": svc.fmt_hours(sum(_realized(r) for r in rows if r.overtime_type == "weekend")),
            "total_hours": svc.fmt_hours(sum(_realized(r) for r in rows)),
            "planned_hours": svc.fmt_hours(sum(r.planned_minutes or 0 for r in rows)),
            "estimated_cost": float(sum(Decimal(r.estimated_cost or 0) for r in rows)),
            "entries": len(rows),
        }

    MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

    def _monthly(rows: list[OvertimeRequest]) -> list[dict]:
        out = []
        for m in range(1, 13):
            sel = [r for r in rows if r.ot_date.month == m]
            out.append({
                "month": m, "month_label": MONTHS[m - 1],
                "weekday_hours": svc.fmt_hours(sum(_realized(r) for r in sel if r.overtime_type == "weekday")),
                "weekend_hours": svc.fmt_hours(sum(_realized(r) for r in sel if r.overtime_type == "weekend")),
                "planned_hours": svc.fmt_hours(sum(r.planned_minutes or 0 for r in sel)),
                "realized_hours": svc.fmt_hours(sum(_realized(r) for r in sel)),
                "estimated_cost": float(sum(Decimal(r.estimated_cost or 0) for r in sel)),
                "entries": len(sel),
            })
        return out

    def _by(rows: list[OvertimeRequest], key_fn, label: str) -> list[dict]:
        buckets: dict[str, dict] = {}
        for r in rows:
            key = key_fn(r) or "(unassigned)"
            b = buckets.setdefault(key, {label: key, "planned": 0, "realized": 0, "weekday": 0, "weekend": 0,
                                        "cost": Decimal(0), "entries": 0, "employees": set()})
            b["planned"] += r.planned_minutes or 0
            b["realized"] += _realized(r)
            b["weekend" if r.overtime_type == "weekend" else "weekday"] += _realized(r)
            b["cost"] += Decimal(r.estimated_cost or 0)
            b["entries"] += 1
            b["employees"].add(r.employee_id)
        out = [{
            label: b[label],
            "planned_hours": svc.fmt_hours(b["planned"]),
            "realized_hours": svc.fmt_hours(b["realized"]),
            "variance_hours": svc.fmt_hours(b["realized"] - b["planned"]),
            "weekday_hours": svc.fmt_hours(b["weekday"]),
            "weekend_hours": svc.fmt_hours(b["weekend"]),
            "estimated_cost": float(b["cost"]),
            "entries": b["entries"],
            "employee_count": len(b["employees"]),
        } for b in buckets.values()]
        return sorted(out, key=lambda x: x["realized_hours"], reverse=True)

    return {
        "year": year,
        "compare_year": compare_year,
        "yearly": [_yearly(rows_cmp, compare_year), _yearly(rows_cur, year)],
        "monthly": _monthly(rows_cur),
        "monthly_compare": _monthly(rows_cmp),
        "by_department": _by(rows_cur, lambda r: r.department, "department"),
        "by_team": _by(rows_cur, lambda r: r.team, "team"),
        "by_employee": _by(rows_cur, lambda r: r.employee_name or r.employee_id, "employee")[:25],
    }


@router.get("/monitoring/exceptions")
async def exception_report(
    year: int = Query(...),
    month: Optional[int] = Query(None, ge=1, le=12),
    threshold_minutes: int = Query(30, ge=0, description="Report rows whose plan/realization gap is at least this large"),
    department: Optional[str] = None,
    team: Optional[str] = None,
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    """Where realization diverged from the plan (site map C.10 "Exception
    Report"), plus two divergences that are easy to miss because they leave
    no gap number at all: an approved order that was never realized, and one
    still stuck in the approval chain long after its date."""
    start = date(year, month or 1, 1)
    end = date(year, 12, 31) if not month else (svc.add_months(start, 1) - timedelta(days=1))

    conditions = [
        OvertimeRequest.ot_date >= start, OvertimeRequest.ot_date <= end,
        OvertimeRequest.plan_status == PLAN_APPROVED,
    ]
    if department:
        conditions.append(OvertimeRequest.department == department)
    if team:
        conditions.append(OvertimeRequest.team == team)
    rows = list((await db.execute(select(OvertimeRequest).where(*conditions))).scalars().all())

    today = date.today()
    exceptions = []
    for r in rows:
        kinds = []
        if r.rz_status == RZ_APPROVED and abs(r.gap_minutes or 0) >= threshold_minutes:
            kinds.append("gap_over_threshold" if (r.gap_minutes or 0) > 0 else "under_realized")
        if r.rz_status in ("pending", "draft") and r.ot_date < today:
            kinds.append("not_realized")
        if r.rz_status in (RZ_SUBMITTED, RZ_TH_APPROVED, RZ_DH_APPROVED) and (today - r.ot_date).days > 7:
            kinds.append("approval_overdue")
        if r.rz_status == RZ_APPROVED and (r.payable_minutes or 0) < (r.actual_minutes or 0):
            kinds.append("hrga_adjusted")
        if not kinds:
            continue
        exceptions.append({
            "id": r.id, "request_no": r.request_no,
            "employee_id": r.employee_id, "employee_name": r.employee_name,
            "department": r.department, "team": r.team,
            "ot_date": r.ot_date.isoformat(), "overtime_type": r.overtime_type,
            "planned_hours": svc.fmt_hours(r.planned_minutes),
            "actual_hours": svc.fmt_hours(r.actual_minutes),
            "payable_hours": svc.fmt_hours(r.payable_minutes),
            "gap_hours": svc.fmt_hours(r.gap_minutes),
            "gap_reason": r.gap_reason,
            "rz_status": r.rz_status,
            "days_open": (today - r.ot_date).days,
            "kinds": kinds,
        })

    exceptions.sort(key=lambda e: abs(e["gap_hours"]), reverse=True)

    def _agg(key: str) -> list[dict]:
        buckets: dict[str, dict] = {}
        for e in exceptions:
            k = e[key] or "(unassigned)"
            b = buckets.setdefault(k, {key: k, "count": 0, "gap_hours": 0.0, "planned_hours": 0.0, "actual_hours": 0.0})
            b["count"] += 1
            b["gap_hours"] = round(b["gap_hours"] + e["gap_hours"], 2)
            b["planned_hours"] = round(b["planned_hours"] + e["planned_hours"], 2)
            b["actual_hours"] = round(b["actual_hours"] + e["actual_hours"], 2)
        return sorted(buckets.values(), key=lambda b: b["count"], reverse=True)

    return {
        "period": {"start": start.isoformat(), "end": end.isoformat(), "year": year, "month": month},
        "threshold_minutes": threshold_minutes,
        "total": len(exceptions),
        "exceptions": exceptions,
        "by_department": _agg("department"),
        "by_team": _agg("team"),
    }


@router.get("/requests")
async def all_requests(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    plan_status: Optional[str] = None,
    rz_status: Optional[str] = None,
    department: Optional[str] = None,
    team: Optional[str] = None,
    search: Optional[str] = None,
    limit: int = Query(500, le=2000),
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    """Company-wide register — HRGA's read-across of every request in any
    state, which none of the role-scoped screens above provide."""
    conditions = []
    if date_from:
        conditions.append(OvertimeRequest.ot_date >= date.fromisoformat(date_from))
    if date_to:
        conditions.append(OvertimeRequest.ot_date <= date.fromisoformat(date_to))
    if plan_status:
        conditions.append(OvertimeRequest.plan_status.in_([s.strip() for s in plan_status.split(",") if s.strip()]))
    if rz_status:
        conditions.append(OvertimeRequest.rz_status.in_([s.strip() for s in rz_status.split(",") if s.strip()]))
    if department:
        conditions.append(OvertimeRequest.department == department)
    if team:
        conditions.append(OvertimeRequest.team == team)
    if search:
        like = f"%{search.strip()}%"
        conditions.append(or_(
            OvertimeRequest.employee_name.ilike(like),
            OvertimeRequest.employee_id.ilike(like),
            OvertimeRequest.request_no.ilike(like),
            OvertimeRequest.task_description.ilike(like),
        ))
    result = await db.execute(
        select(OvertimeRequest).where(*conditions)
        .order_by(OvertimeRequest.ot_date.desc(), OvertimeRequest.id.desc()).limit(limit)
    )
    return {"requests": [svc.serialize_request(r) for r in result.scalars().all()]}


# ─────────────────────────────────────────
# D.1 — APPROVAL MATRIX
# ─────────────────────────────────────────

@router.get("/approval-matrix")
async def list_matrix(
    search: Optional[str] = None,
    department: Optional[str] = None,
    role_level: Optional[str] = None,
    unassigned_only: bool = False,
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    """Every active employee joined to their matrix row — employees with no
    row are listed too, with role_level "member" and no approvers, because
    the point of this screen is finding exactly those."""
    emp_result = await db.execute(
        select(Employee).where(Employee.employment_status == "Active").order_by(Employee.department, Employee.team, Employee.full_name)
    )
    employees = list(emp_result.scalars().all())
    names = {e.user_id: e.full_name for e in employees}

    matrix_rows = {m.employee_id: m for m in (await db.execute(select(OvertimeApprovalMatrix))).scalars().all()}
    # Names for approvers who may themselves be inactive/resigned.
    missing = {m.team_head_id for m in matrix_rows.values()} | {m.dept_head_id for m in matrix_rows.values()}
    missing = {i for i in missing if i and i not in names}
    if missing:
        extra = await db.execute(select(Employee).where(Employee.user_id.in_(missing)))
        names.update({e.user_id: e.full_name for e in extra.scalars().all()})

    out = []
    for e in employees:
        m = matrix_rows.get(e.user_id)
        row = {
            "id": m.id if m else None,
            "employee_id": e.user_id,
            "employee_name": e.full_name,
            "department": e.department,
            "team": e.team,
            "job_title": e.job_title,
            "role_level": m.role_level if m else ROLE_MEMBER,
            "team_head_id": m.team_head_id if m else None,
            "team_head_name": names.get(m.team_head_id) if m else None,
            "dept_head_id": m.dept_head_id if m else None,
            "dept_head_name": names.get(m.dept_head_id) if m else None,
            "is_active": m.is_active if m else False,
            "configured": m is not None,
            "updated_by": m.updated_by if m else None,
            "updated_at": m.updated_at.isoformat() if m and m.updated_at else None,
        }
        if unassigned_only and row["configured"] and row["team_head_id"] and row["dept_head_id"]:
            continue
        if department and (e.department or "") != department:
            continue
        if role_level and row["role_level"] != role_level:
            continue
        if search:
            q = search.strip().lower()
            if q not in (e.full_name or "").lower() and q not in (e.user_id or "").lower():
                continue
        out.append(row)

    return {
        "rows": out,
        "role_levels": list(_ROLE_LEVELS),
        "summary": {
            "total_employees": len(employees),
            "configured": sum(1 for r in out if r["configured"]),
            "unconfigured": sum(1 for r in out if not r["configured"]),
        },
    }


def _validate_matrix(employee_id: str, role_level: str, team_head_id: Optional[str], dept_head_id: Optional[str]) -> None:
    if role_level not in _ROLE_LEVELS:
        raise HTTPException(400, f"role_level must be one of {list(_ROLE_LEVELS)}")
    if team_head_id and team_head_id == employee_id:
        raise HTTPException(400, "An employee cannot be their own Team Head.")
    if dept_head_id and dept_head_id == employee_id:
        raise HTTPException(400, "An employee cannot be their own Department Head.")
    if team_head_id and dept_head_id and team_head_id == dept_head_id:
        raise HTTPException(400, "Team Head and Department Head must be two different people — a request would otherwise be approved twice by the same person.")


@router.post("/approval-matrix")
async def upsert_matrix(
    body: MatrixIn,
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    emp = (await db.execute(select(Employee).where(Employee.user_id == body.employee_id))).scalar_one_or_none()
    if not emp:
        raise HTTPException(404, f"Employee {body.employee_id} not found")
    _validate_matrix(body.employee_id, body.role_level, body.team_head_id, body.dept_head_id)

    row = (await db.execute(
        select(OvertimeApprovalMatrix).where(OvertimeApprovalMatrix.employee_id == body.employee_id)
    )).scalar_one_or_none()
    if not row:
        row = OvertimeApprovalMatrix(employee_id=body.employee_id)
        db.add(row)
    row.employee_name = emp.full_name
    row.department = emp.department
    row.team = emp.team
    row.role_level = body.role_level
    row.team_head_id = body.team_head_id or None
    row.dept_head_id = body.dept_head_id or None
    row.is_active = body.is_active
    row.updated_by = actor.email
    await db.flush()
    return {"id": row.id, "employee_id": row.employee_id, "message": "Approval matrix saved"}


@router.post("/approval-matrix/bulk")
async def bulk_matrix(
    body: MatrixBulkIn,
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    """Apply the same heads (and/or role) to a whole team at once — the
    realistic way this table gets filled the first time."""
    if not body.employee_ids:
        raise HTTPException(400, "Select at least one employee.")
    if body.role_level and body.role_level not in _ROLE_LEVELS:
        raise HTTPException(400, f"role_level must be one of {list(_ROLE_LEVELS)}")

    emps = {e.user_id: e for e in (await db.execute(
        select(Employee).where(Employee.user_id.in_(body.employee_ids)))).scalars().all()}
    existing = {m.employee_id: m for m in (await db.execute(
        select(OvertimeApprovalMatrix).where(OvertimeApprovalMatrix.employee_id.in_(body.employee_ids)))).scalars().all()}

    applied, skipped = 0, []
    for emp_id in body.employee_ids:
        emp = emps.get(emp_id)
        if not emp:
            skipped.append({"employee_id": emp_id, "reason": "not found"})
            continue
        row = existing.get(emp_id)
        th = body.team_head_id if body.team_head_id is not None else (row.team_head_id if row else None)
        dh = body.dept_head_id if body.dept_head_id is not None else (row.dept_head_id if row else None)
        level = body.role_level or (row.role_level if row else ROLE_MEMBER)
        try:
            _validate_matrix(emp_id, level, th, dh)
        except HTTPException as exc:
            # One bad member (e.g. the team head is in their own selection)
            # must not sink the other 20 — skip and report it back.
            skipped.append({"employee_id": emp_id, "reason": exc.detail})
            continue
        if not row:
            row = OvertimeApprovalMatrix(employee_id=emp_id)
            db.add(row)
        row.employee_name, row.department, row.team = emp.full_name, emp.department, emp.team
        row.role_level, row.team_head_id, row.dept_head_id = level, th or None, dh or None
        row.is_active = True
        row.updated_by = actor.email
        applied += 1

    await db.flush()
    return {"applied": applied, "skipped": skipped, "message": f"{applied} employee(s) updated"}


@router.delete("/approval-matrix/{matrix_id}")
async def delete_matrix(
    matrix_id: int,
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    row = (await db.execute(select(OvertimeApprovalMatrix).where(OvertimeApprovalMatrix.id == matrix_id))).scalar_one_or_none()
    if not row:
        raise HTTPException(404, "Approval matrix row not found")
    await db.delete(row)
    return {"message": "Approval matrix row removed"}


@router.get("/approver-options")
async def approver_options(
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    """Candidate approvers for the Team/Dept Head pickers — every active
    employee, with their matrix role shown so HRGA can see at a glance who
    is already designated a head."""
    result = await db.execute(
        select(Employee).where(Employee.employment_status == "Active").order_by(Employee.full_name)
    )
    employees = list(result.scalars().all())
    levels = {m.employee_id: m.role_level for m in (await db.execute(select(OvertimeApprovalMatrix))).scalars().all()}
    return [
        {"employee_id": e.user_id, "full_name": e.full_name, "department": e.department,
         "team": e.team, "job_title": e.job_title, "role_level": levels.get(e.user_id, ROLE_MEMBER)}
        for e in employees
    ]


# ─────────────────────────────────────────
# D.3 — OVERTIME RULES
# ─────────────────────────────────────────

@router.get("/rules")
async def admin_list_rules(
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(select(OvertimeRule).order_by(OvertimeRule.seq, OvertimeRule.id))
    return [{"id": r.id, "seq": r.seq, "title": r.title, "content": r.content, "is_active": r.is_active,
             "updated_by": r.updated_by, "updated_at": r.updated_at.isoformat() if r.updated_at else None}
            for r in result.scalars().all()]


@router.post("/rules")
async def create_rule(
    body: RuleIn,
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    if not body.title.strip():
        raise HTTPException(400, "Rule title is required.")
    seq = body.seq
    if not seq:
        seq = ((await db.execute(select(func.coalesce(func.max(OvertimeRule.seq), 0)))).scalar() or 0) + 1
    row = OvertimeRule(seq=seq, title=body.title.strip(), content=body.content,
                       is_active=body.is_active, updated_by=actor.email)
    db.add(row)
    await db.flush()
    return {"id": row.id, "message": "Rule added"}


@router.put("/rules/{rule_id}")
async def update_rule(
    rule_id: int,
    body: RuleIn,
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    row = (await db.execute(select(OvertimeRule).where(OvertimeRule.id == rule_id))).scalar_one_or_none()
    if not row:
        raise HTTPException(404, "Rule not found")
    row.title, row.content, row.seq, row.is_active = body.title.strip(), body.content, body.seq, body.is_active
    row.updated_by = actor.email
    return {"message": "Rule updated"}


@router.delete("/rules/{rule_id}")
async def delete_rule(
    rule_id: int,
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    row = (await db.execute(select(OvertimeRule).where(OvertimeRule.id == rule_id))).scalar_one_or_none()
    if not row:
        raise HTTPException(404, "Rule not found")
    await db.delete(row)
    return {"message": "Rule deleted"}


# ─────────────────────────────────────────
# D.4 — CUT-OFF CONFIGURATION
# ─────────────────────────────────────────

@router.get("/cutoff")
async def get_cutoff(
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    cfg = await svc.get_active_cutoff(db)
    today = date.today()
    anchor = today if today.day <= cfg.end_day else svc.add_months(today, 1)
    current = svc.cutoff_period_for_key(f"{anchor.year:04d}-{anchor.month:02d}", cfg.start_day, cfg.end_day)
    history = (await db.execute(
        select(OvertimeCutoffConfig).order_by(OvertimeCutoffConfig.id.desc()).limit(10))).scalars().all()
    return {
        "id": cfg.id, "name": cfg.name, "start_day": cfg.start_day, "end_day": cfg.end_day,
        "updated_by": cfg.updated_by, "updated_at": cfg.updated_at.isoformat() if cfg.updated_at else None,
        "current_period": {"key": current.key, "label": current.label,
                           "start": current.start.isoformat(), "end": current.end.isoformat()},
        "history": [{"id": h.id, "name": h.name, "start_day": h.start_day, "end_day": h.end_day,
                     "is_active": h.is_active, "updated_by": h.updated_by,
                     "updated_at": h.updated_at.isoformat() if h.updated_at else None} for h in history],
    }


@router.put("/cutoff")
async def update_cutoff(
    body: CutoffIn,
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    """Supersede the active window with a new row rather than editing the
    old one in place — a calculation re-run for a closed period must still
    be able to show the window that was in force when it was closed."""
    if not (1 <= body.start_day <= 28) or not (1 <= body.end_day <= 28):
        raise HTTPException(400, "Cut-off days must be between 1 and 28 — a later day does not exist in every month.")
    if body.start_day == body.end_day:
        raise HTTPException(400, "Start and finish day cannot be the same.")

    current = await svc.get_active_cutoff(db)
    if current.start_day == body.start_day and current.end_day == body.end_day and current.name == body.name:
        return {"message": "No change — cut-off already set to these dates.", "id": current.id}

    current.is_active = False
    row = OvertimeCutoffConfig(name=body.name, start_day=body.start_day, end_day=body.end_day,
                               is_active=True, updated_by=actor.email)
    db.add(row)
    await db.flush()
    return {"id": row.id, "message": "Cut-off configuration updated"}


# ─────────────────────────────────────────
# D.5 — NOTIFICATION SETTINGS
# ─────────────────────────────────────────

@router.get("/notification-settings")
async def get_notification_settings(
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    rows = await svc.ensure_notification_settings(db)
    return [{"event_key": r.event_key, "label": r.label, "in_app": r.in_app, "email": r.email,
             "is_active": r.is_active, "updated_by": r.updated_by,
             "updated_at": r.updated_at.isoformat() if r.updated_at else None} for r in rows]


@router.put("/notification-settings")
async def update_notification_settings(
    body: list[NotificationSettingIn],
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    rows = {r.event_key: r for r in await svc.ensure_notification_settings(db)}
    for item in body:
        row = rows.get(item.event_key)
        if not row:
            raise HTTPException(400, f"Unknown notification event '{item.event_key}'")
        row.in_app, row.email, row.is_active = item.in_app, item.email, item.is_active
        row.updated_by = actor.email
    return {"message": "Notification settings saved"}


# ─────────────────────────────────────────
# OVERTIME RATES (cost estimation)
# ─────────────────────────────────────────

@router.get("/rates")
async def list_rates(
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(select(OvertimeRateConfig).order_by(OvertimeRateConfig.grade))
    rates = [{"id": r.id, "grade": r.grade, "hourly_rate": float(r.hourly_rate or 0),
              "updated_by": r.updated_by, "updated_at": r.updated_at.isoformat() if r.updated_at else None}
             for r in result.scalars().all()]
    grades = (await db.execute(
        select(Employee.employee_grade).where(Employee.employee_grade.is_not(None)).distinct()
    )).all()
    known = sorted({(g[0] or "").strip() for g in grades if (g[0] or "").strip()})
    return {"rates": rates, "employee_grades": known}


@router.post("/rates")
async def upsert_rate(
    body: RateIn,
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    grade = body.grade.strip()
    if not grade:
        raise HTTPException(400, "Grade is required. Use 'DEFAULT' for the fallback rate.")
    if body.hourly_rate < 0:
        raise HTTPException(400, "Hourly rate cannot be negative.")
    row = (await db.execute(select(OvertimeRateConfig).where(OvertimeRateConfig.grade.ilike(grade)))).scalar_one_or_none()
    if not row:
        row = OvertimeRateConfig(grade=grade)
        db.add(row)
    row.hourly_rate = Decimal(str(body.hourly_rate))
    row.updated_by = actor.email
    await db.flush()
    return {"id": row.id, "message": f"Rate for grade '{grade}' saved"}


@router.delete("/rates/{rate_id}")
async def delete_rate(
    rate_id: int,
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    row = (await db.execute(select(OvertimeRateConfig).where(OvertimeRateConfig.id == rate_id))).scalar_one_or_none()
    if not row:
        raise HTTPException(404, "Rate not found")
    await db.delete(row)
    return {"message": "Rate deleted"}


# ─────────────────────────────────────────
# LOOKUPS
# ─────────────────────────────────────────

@router.get("/lov")
async def lookups(
    actor: OvertimeActor = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    depts = (await db.execute(select(OvertimeRequest.department).distinct())).all()
    teams = (await db.execute(select(OvertimeRequest.team).distinct())).all()
    emp_depts = (await db.execute(
        select(Employee.department).where(Employee.employment_status == "Active").distinct())).all()
    emp_teams = (await db.execute(
        select(Employee.team).where(Employee.employment_status == "Active").distinct())).all()
    clean = lambda rows: sorted({(r[0] or "").strip() for r in rows if (r[0] or "").strip()})
    return {
        "departments": clean(list(depts) + list(emp_depts)),
        "teams": clean(list(teams) + list(emp_teams)),
    }
