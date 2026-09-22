"""
AP VAT In Listing Report Router
Route prefix : /api/v1/dashboard/accounting/ap-vat-in
Required role: accounting_staff OR admin

Endpoints:
  GET /              — JSON rows for the on-screen table
  GET /export        — Excel download, laid out to match
                        sumber/CUSTOM-ADDITIONAL COA DESCPT IN VAT MONTHLY REPORT.xlsx
"""
import io
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse

from app.dependencies import CurrentUser, Roles, require_role
from app.services.ap_vat_in_service import COLUMNS, _TOTAL_FIELDS, get_ap_vat_in_listing

router = APIRouter()


@router.get("")
async def get_ap_vat_in(
    period: str = Query(..., description="Oracle GL period, e.g. AUG-26"),
    user: CurrentUser = Depends(require_role(Roles.ACCOUNTING)),
):
    return await get_ap_vat_in_listing(period)


@router.get("/export")
async def export_ap_vat_in(
    period: str = Query(..., description="Oracle GL period, e.g. AUG-26"),
    user: CurrentUser = Depends(require_role(Roles.ACCOUNTING)),
):
    result = await get_ap_vat_in_listing(period)
    if not result.get("success"):
        raise HTTPException(400, result.get("error") or "Failed to load AP VAT In Listing data")
    try:
        return _build_ap_vat_in_xlsx(result, period)
    except Exception as e:
        raise HTTPException(500, f"Excel generation failed: {e}")


def _build_ap_vat_in_xlsx(result: dict, period: str) -> StreamingResponse:
    rows = result["data"]
    totals = result["totals"]
    n_cols = len(COLUMNS)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AP VAT In Listing Report"

    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    fillHdr = PatternFill("solid", fgColor="D9E1F2")

    def merge(r1, c1, r2, c2, value=None, style=None):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        cell = ws.cell(row=r1, column=c1, value=value)
        if style:
            cell.font, cell.alignment = style
        return cell

    merge(1, 1, 1, n_cols, "PT CKD OTTO Pharmaceuticals", (Font(bold=True, size=13), center))
    merge(2, 1, 2, n_cols, "VAT In Listing Report", (bold, center))
    merge(3, 1, 3, n_cols, f"Period : {period.upper()}", (bold, center))

    # "ADDITIONAL" group label spanning the 3 COA columns, matching the
    # source template's merged M4:O4 header.
    coa_start = next(i for i, (k, _) in enumerate(COLUMNS, start=1) if k == "coa_full")
    coa_end = next(i for i, (k, _) in enumerate(COLUMNS, start=1) if k == "coa_description")
    merge(4, coa_start, 4, coa_end, "ADDITIONAL", (bold, center))
    ws.cell(row=4, column=coa_start).fill = fillHdr

    HEADER_ROW = 5
    for i, (_, label) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=HEADER_ROW, column=i, value=label)
        c.font, c.alignment = bold, center
        c.fill = fillHdr

    DATE_FIELDS = {"posted_date", "supplier_tax_invoice_date"}
    NUMFMT = "#,##0.####"

    r_idx = HEADER_ROW + 1
    for row in rows:
        for i, (key, _) in enumerate(COLUMNS, start=1):
            v = row.get(key)
            if key in DATE_FIELDS and v:
                try:
                    v = datetime.fromisoformat(v)
                except (TypeError, ValueError):
                    pass
            cell = ws.cell(row=r_idx, column=i, value=v)
            if key in DATE_FIELDS and v:
                cell.number_format = "DD-MON-YY"
            elif key in _TOTAL_FIELDS:
                cell.number_format = NUMFMT
        r_idx += 1

    # Total row
    label_col = next(i for i, (k, _) in enumerate(COLUMNS, start=1) if k == "supplier_name")
    ws.cell(row=r_idx, column=label_col, value=f"Total VAT In {period.upper()} : ").font = bold
    for i, (key, _) in enumerate(COLUMNS, start=1):
        if key in _TOTAL_FIELDS:
            c = ws.cell(row=r_idx, column=i, value=totals.get(key, 0))
            c.font = bold
            c.number_format = NUMFMT

    ws.column_dimensions["A"].width = 18
    for i, (key, _) in enumerate(COLUMNS, start=1):
        letter = get_column_letter(i)
        if key == "description":
            ws.column_dimensions[letter].width = 40
        elif key == "supplier_name":
            ws.column_dimensions[letter].width = 28
        elif key in ("coa_full", "coa_description"):
            ws.column_dimensions[letter].width = 22
        else:
            ws.column_dimensions[letter].width = 15
    ws.freeze_panes = f"A{HEADER_ROW + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"ap_vat_in_listing_{period.upper()}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"},
    )
