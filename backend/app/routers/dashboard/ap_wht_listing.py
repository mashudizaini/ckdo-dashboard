"""
AP Withholding Tax Listing Report Router
Route prefix : /api/v1/dashboard/accounting/ap-wht-listing
Required role: accounting_staff OR admin

Endpoints:
  GET /              — JSON rows for the on-screen table
  GET /export        — Excel download, laid out to match
                        sumber/Witholding Tax Listing - Output.xlsx
"""
import io
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse

from app.dependencies import CurrentUser, Roles, require_role
from app.services.ap_wht_listing_service import COLUMNS, _TOTAL_FIELDS, get_ap_wht_listing

router = APIRouter()


@router.get("")
async def get_ap_wht(
    period: str = Query(..., description="Oracle GL period, e.g. JAN-26"),
    user: CurrentUser = Depends(require_role(Roles.ACCOUNTING)),
):
    return await get_ap_wht_listing(period)


@router.get("/export")
async def export_ap_wht(
    period: str = Query(..., description="Oracle GL period, e.g. JAN-26"),
    user: CurrentUser = Depends(require_role(Roles.ACCOUNTING)),
):
    result = await get_ap_wht_listing(period)
    if not result.get("success"):
        raise HTTPException(400, result.get("error") or "Failed to load AP Withholding Tax Listing data")
    try:
        return _build_ap_wht_xlsx(result, period)
    except Exception as e:
        raise HTTPException(500, f"Excel generation failed: {e}")


def _build_ap_wht_xlsx(result: dict, period: str) -> StreamingResponse:
    rows = result["data"]
    totals = result["totals"]
    n_cols = len(COLUMNS)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AP Withholding Tax Listing"

    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    fillHdr = PatternFill("solid", fgColor="D9E1F2")

    def merge(r1, c1, r2, c2, value=None, style=None):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        cell = ws.cell(row=r1, column=c1, value=value)
        if style:
            cell.font, cell.alignment = style
        return cell

    merge(1, 1, 1, n_cols, "PT CKD OTTO Pharmaceuticals", (Font(bold=True, size=13), center))
    merge(2, 1, 2, n_cols, "Withholding Tax Listing", (bold, center))
    merge(3, 1, 3, n_cols, f"Period : {period.upper()}", (bold, center))

    HEADER_ROW = 5
    for i, (_, label) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=HEADER_ROW, column=i, value=label)
        c.font, c.alignment = bold, center
        c.fill = fillHdr

    DATE_FIELDS = {"invoice_date", "supplier_tax_invoice_date", "payment_date"}
    NUMFMT = "#,##0.####"

    r_idx = HEADER_ROW + 1
    for row in rows:
        for i, (key, _) in enumerate(COLUMNS, start=1):
            v = row.get(key)
            if key in DATE_FIELDS and v:
                try:
                    v = datetime.fromisoformat(v)
                except (TypeError, ValueError):
                    pass
            cell = ws.cell(row=r_idx, column=i, value=v)
            if key in DATE_FIELDS and v:
                cell.number_format = "DD-MON-YY"
            elif key in _TOTAL_FIELDS or key == "tax_rate":
                cell.number_format = NUMFMT
        r_idx += 1

    # Total row — label spans A:L (up through Tax Rate), sums in Gross/WHT
    # Amount, matching the source template's merged A75:L75 + O75:P75.
    label_end_col = next(i for i, (k, _) in enumerate(COLUMNS, start=1) if k == "tax_rate")
    merge(r_idx, 1, r_idx, label_end_col, f"Total WHT {period.upper()} : ", (bold, None))
    tail_start_col = next(i for i, (k, _) in enumerate(COLUMNS, start=1) if k == "description")
    ws.merge_cells(start_row=r_idx, start_column=tail_start_col, end_row=r_idx, end_column=n_cols)
    for i, (key, _) in enumerate(COLUMNS, start=1):
        if key in _TOTAL_FIELDS:
            c = ws.cell(row=r_idx, column=i, value=totals.get(key, 0))
            c.font = bold
            c.number_format = NUMFMT

    ws.column_dimensions["A"].width = 6
    for i, (key, _) in enumerate(COLUMNS, start=1):
        letter = get_column_letter(i)
        if key in ("description", "supplier_name"):
            ws.column_dimensions[letter].width = 32
        elif key == "coa_description":
            ws.column_dimensions[letter].width = 24
        elif key in ("invoice_no", "coa_full", "npwp", "faktur_pajak"):
            ws.column_dimensions[letter].width = 20
        elif key == "nomor":
            pass
        else:
            ws.column_dimensions[letter].width = 15
    ws.freeze_panes = f"A{HEADER_ROW + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"ap_wht_listing_{period.upper()}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"},
    )
